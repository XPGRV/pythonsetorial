import os
import re
import copy
import datetime as dt
import requests
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DRIVE_ROOT = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
)
XLSX_NAME = "FrangoBR.xlsm"
SHEET_NAME = "Processados"


def localizar_planilha() -> Path:
    """Localiza a planilha sem depender dos nomes das pastas intermediárias."""
    encontrados = list(DRIVE_ROOT.rglob(XLSX_NAME))

    if not encontrados:
        raise FileNotFoundError(
            f"Planilha {XLSX_NAME!r} não encontrada dentro de {DRIVE_ROOT}"
        )

    if len(encontrados) > 1:
        caminhos = "\n".join(f"  - {path}" for path in encontrados)
        raise RuntimeError(
            f"Foram encontradas várias planilhas chamadas {XLSX_NAME!r}:\n"
            f"{caminhos}"
        )

    return encontrados[0]


# Baixa a planilha do Google Drive (em vez do drive montado em G:\).
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import drive_io
import ibge_sidra
_fid = drive_io.find_file_id(XLSX_NAME)
EXCEL_PATH = drive_io.download_to_temp(_fid)

# Colunas no Excel (1-based): C=3 D=4 E=5 F=6 G=7 H=8
COLS = {
    "Presunto":  3,
    "Salsicha":  4,
    "Linguica":  5,
    "Mortadela": 6,
    "Salame":    7,
    "Margarina": 8,
}

# Mapeamento código IBGE → nome interno, por agregado
# Cada período usa POF diferente, logo os códigos mudam
AGGREGATES = [
    {
        "id": 655,
        "label": "Jan/2000–Jun/2006",
        "periodos": "|".join(
            f"{y}{m:02d}"
            for y in range(2000, 2007)
            for m in range(1, 13)
            if (y, m) <= (2006, 6)
        ),
        "codigos": {
            "7336": "Presunto",
            "7338": "Salsicha",
            "7339": "Linguica",
            "7341": "Mortadela",
            "7342": "Salame",
            "7387": "Margarina",
        },
    },
    {
        "id": 2938,
        "label": "Jul/2006–Dez/2011",
        "periodos": "|".join(
            f"{y}{m:02d}"
            for y in range(2006, 2012)
            for m in range(1, 13)
            if (y, m) >= (2006, 7)
        ),
        "codigos": {
            "7336": "Presunto",
            "7338": "Salsicha",
            "7339": "Linguica",
            "7341": "Mortadela",
            # Salame ausente neste período (não faz parte da cesta POF 2002-2003)
            "7387": "Margarina",
        },
    },
    {
        "id": 1419,
        "label": "Jan/2012–Dez/2019",
        "periodos": "|".join(
            f"{y}{m:02d}"
            for y in range(2012, 2020)
            for m in range(1, 13)
        ),
        "codigos": {
            "7336":  "Presunto",
            "12304": "Salsicha",
            "7339":  "Linguica",
            "7341":  "Mortadela",
            "12305": "Salame",
            "12395": "Margarina",
        },
    },
    {
        "id": 7060,
        "label": "Jan/2020–presente (dinâmico)",
        "periodos": None,  # será buscado dinamicamente via metadados SIDRA
        "inicio": "202001",  # usado se a descoberta de períodos não responder
        "codigos": {
            "7336":  "Presunto",
            "12304": "Salsicha",
            "7339":  "Linguica",
            "7341":  "Mortadela",
            "12305": "Salame",
            "12395": "Margarina",
        },
    },
]

def meses_ate_hoje(inicio: str) -> list[str]:
    """['AAAAMM', ...] do mês inicial até o mês corrente."""
    ano, mes = int(inicio[:4]), int(inicio[4:])
    hoje = dt.date.today()
    saida = []
    while (ano, mes) <= (hoje.year, hoje.month):
        saida.append(f"{ano}{mes:02d}")
        mes += 1
        if mes > 12:
            ano, mes = ano + 1, 1
    return saida


def get_periodos_disponiveis(agg: dict) -> list[str]:
    """Períodos a consultar para um agregado.

    A lista publicada só existe no servicodados; o endereço antigo não tem
    equivalente. Se ele não responder, o calendário serve: mês sem publicação
    não volta na consulta de dados.
    """
    if agg["periodos"]:
        return agg["periodos"].split("|")

    url = f"https://servicodados.ibge.gov.br/api/v3/agregados/{agg['id']}/periodos"
    try:
        periodos = [p["id"] for p in ibge_sidra.buscar_json(url, timeout=30)]
        print(f"    -> {len(periodos)} periodos disponiveis ({periodos[0]} a {periodos[-1]})")
        return periodos
    except Exception as exc:
        periodos = meses_ate_hoje(agg.get("inicio", "202001"))
        print(f"    -> /periodos indisponivel ({type(exc).__name__}); "
              f"usando o calendario, {len(periodos)} meses.")
        return periodos

def row_for(periodo: str) -> int:
    """Converte 'AAAAMM' para linha do Excel. Linha 3 = jan/2000."""
    y, m = int(periodo[:4]), int(periodo[4:])
    return 3 + (y - 2000) * 12 + (m - 1)

def fetch_aggregate(agg: dict) -> dict:
    """Retorna {periodo: {nome: valor}} para um agregado."""
    codigos = agg["codigos"]

    if agg["periodos"] is None:
        print(f"\n  Consultando períodos disponíveis para tabela {agg['id']}...")
    periodos = get_periodos_disponiveis(agg)

    print(f"\n  Buscando {agg['label']}...")
    # Pelos dois endereços do IBGE. Saída plana: D3C período, D4C categoria da
    # classificação 315, V valor.
    linhas = ibge_sidra.consultar(
        tabela=str(agg["id"]),
        variaveis=["63"],
        periodos=periodos,
        classificacoes=[("315", list(codigos))],
        timeout=60,
    )

    result = {}  # {periodo: {nome: float}}
    for linha in linhas:
        nome = codigos.get(str(linha.get("D4C")))
        if not nome:
            continue
        periodo = str(linha.get("D3C") or "")
        if len(periodo) != 6:
            continue
        try:
            valor = float(str(linha.get("V")).replace(",", "."))
        except (ValueError, TypeError):
            valor = None
        result.setdefault(periodo, {})[nome] = valor

    return result

# ── Coleta de todos os agregados ─────────────────────────────────────────────
print("=== Coletando dados da API do IBGE ===")
all_data = {}  # {periodo: {nome: valor}}

for agg in AGGREGATES:
    agg_data = fetch_aggregate(agg)
    for periodo, valores in agg_data.items():
        all_data.setdefault(periodo, {}).update(valores)
    print(f"  {len(agg_data)} períodos retornados")

# ── Verificação de gaps ───────────────────────────────────────────────────────
print("\n=== Verificando gaps ===")
nomes = list(COLS.keys())
gaps = []

for periodo in sorted(all_data.keys()):
    for nome in nomes:
        val = all_data[periodo].get(nome)
        if val is None:
            gaps.append((periodo, nome))

if gaps:
    print(f"  {len(gaps)} valores ausentes encontrados:")
    for periodo, nome in gaps:
        print(f"    {periodo[:4]}/{periodo[4:]} -> {nome}")
else:
    print("  Nenhum gap encontrado.")

# Verifica meses esperados que não vieram na API
print("\n=== Verificando meses esperados vs. recebidos ===")
expected = set(
    f"{y}{m:02d}"
    for y in range(2000, 2027)
    for m in range(1, 13)
    if (y, m) <= (2026, 3)
)
received = set(all_data.keys())
missing_periods = sorted(expected - received)
if missing_periods:
    print(f"  {len(missing_periods)} meses esperados não retornados pela API:")
    for p in missing_periods:
        print(f"    {p[:4]}/{p[4:]}")
else:
    print("  Todos os meses esperados foram retornados.")

# ── Helpers openpyxl ─────────────────────────────────────────────────────────

def copy_cell_style(src, dst):
    """Copia formatação completa de uma célula para outra."""
    if src.has_style:
        dst.font      = copy.copy(src.font)
        dst.border    = copy.copy(src.border)
        dst.fill      = copy.copy(src.fill)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format

def adjust_formula(formula: str, row_offset: int) -> str:
    """Ajusta referências de linha em fórmulas (ex: A3 → A4). Preserva $-absolutos."""
    if not formula or not str(formula).startswith("="):
        return formula
    def replace(m):
        col_abs, col, row_abs, row = m.group(1), m.group(2), m.group(3), m.group(4)
        if row_abs == "$":          # linha absoluta → não mexe
            return m.group(0)
        return f"{col_abs}{col}{row_abs}{int(row) + row_offset}"
    return re.sub(r"(\$?)([A-Za-z]+)(\$?)(\d+)", replace, formula)

# ── Escrita no Excel via openpyxl (sem abrir Excel) ──────────────────────────
print("\n=== Escrevendo no Excel (sem abrir Excel) ===")

# Avisa se o arquivo estiver travado pelo Excel
lock_path = os.path.join(os.path.dirname(EXCEL_PATH), "~$" + os.path.basename(EXCEL_PATH))
if os.path.exists(lock_path):
    # Tenta abrir exclusivamente — se conseguir, o lock é orfao e pode ser removido
    try:
        with open(lock_path, "r+b"):
            pass
        os.remove(lock_path)
        print("  Lock orfao removido, continuando...")
    except (PermissionError, OSError):
        print("  AVISO: arquivo aberto no Excel. Feche-o antes de rodar o script.")
        print("  Abortando para não corromper o arquivo.")
        raise SystemExit(1)

wb = load_workbook(EXCEL_PATH, keep_vba=True)

# "Processados" pode estar muito oculta → acessa via _sheets
ws = None
for sheet in wb._sheets:
    if sheet.title == SHEET_NAME:
        ws = sheet
        break
if ws is None:
    raise ValueError(f"Aba '{SHEET_NAME}' não encontrada no arquivo.")

ROW_START = 3   # jan/2000
ROW_END   = row_for(max(all_data.keys()))
COL_START = 3   # coluna C
COL_END   = 8   # coluna H
FORMULA_COL_START = 10  # coluna J
FORMULA_COL_END   = 16  # coluna P

# Detecta última linha com dado antes desta execução
old_last_row = ROW_START - 1
for r in range(ROW_END, ROW_START - 1, -1):
    if any(ws.cell(r, c).value is not None for c in range(COL_START, COL_END + 1)):
        old_last_row = r
        break

# Escreve valores e aplica formato 0.00%
written = 0
skipped = 0
for periodo in sorted(all_data.keys()):
    row = row_for(periodo)
    if row < ROW_START or row > ROW_END:
        continue
    for nome, col in COLS.items():
        val = all_data[periodo].get(nome)
        cell = ws.cell(row, col)
        if val is not None:
            cell.value = val / 100
            cell.number_format = "0.00%"
            written += 1
        else:
            skipped += 1

print(f"  {written} valores escritos.")

# Para cada linha nova: copia formatação de C-H e fórmulas+formato de J-P
if ROW_END > old_last_row:
    n_new = ROW_END - old_last_row
    print(f"  {n_new} linha(s) nova(s) — copiando formatação e fórmulas...")
    for new_row in range(old_last_row + 1, ROW_END + 1):
        prev_row = new_row - 1
        offset   = new_row - prev_row  # sempre 1

        # Formatação C-H (não sobrescreve valor, só estilo)
        for col in range(COL_START, COL_END + 1):
            copy_cell_style(ws.cell(prev_row, col), ws.cell(new_row, col))
            ws.cell(new_row, col).number_format = "0.00%"  # garante formato decimal

        # Fórmulas + formatação J-P
        for col in range(FORMULA_COL_START, FORMULA_COL_END + 1):
            src = ws.cell(prev_row, col)
            dst = ws.cell(new_row,  col)
            copy_cell_style(src, dst)
            dst.value = adjust_formula(src.value, offset)

wb.save(EXCEL_PATH)
drive_io.update_file(_fid, EXCEL_PATH)   # sobe de volta pro Drive (in-place)
print("  Arquivo salvo.")

print(f"  Concluído! {written} células escritas, {skipped} ausentes (gaps).")
print("\nResumo de gaps por subitem:")
gap_count = {}
for _, nome in gaps:
    gap_count[nome] = gap_count.get(nome, 0) + 1
for nome, count in sorted(gap_count.items()):
    print(f"  {nome}: {count} meses sem dado")
