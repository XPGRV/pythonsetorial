"""
atualizar_frangobr.py  (versao openpyxl / nuvem)

Consolidador do FrangoBR — "puxa" as formulas-template para baixo, ate as
linhas novas, em duas abas:

Aba BBG_Dados:
  Colunas J, K, L alinhadas com a ultima linha com dado em cols E->I (ref col E).

Aba FrangoBR:
  (i)   Mes corrente:  D->I, L, O->R, T  — ate o mes corrente
  (ii)  SIF:           W                 — 2 meses antes do ultimo dado em Production col F
  (iii) IBGE:          X, Y              — ate o ultimo dado em Production col C
  (iv)  APINCO:        Z                 — ate o ultimo dado em Production col J

Mapeamento de linhas:
  FrangoBR:   linha 5 = jan/2000  ->  row = 5 + (ano-2000)*12 + (mes-1)
  Production: linha 3 = jan/2000  ->  FrangoBR_row = Production_row + 2

Le/grava o FrangoBR.xlsm via Google Drive (drive_io). Roda na nuvem, sem Excel.
As formulas ganham valor quando o Excel (PC Bloomberg) recalcular.
Requer: pip install openpyxl google-api-python-client google-auth
"""

import re
import sys
import datetime
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import drive_io

XLSX_NAME = "FrangoBR.xlsm"


# ── Deslocamento de formula ("puxar pra baixo") ───────────────────────────────
_REF = re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?![A-Za-z_(])")


def shift_formula(formula, delta: int):
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(m):
        col_abs, col, row_abs, row = m.groups()
        if row_abs == "$":
            return m.group(0)
        return f"{col_abs}{col}{row_abs}{int(row) + delta}"

    return _REF.sub(repl, formula)


# ── Helpers ───────────────────────────────────────────────────────────────────

def col_letter(n: int) -> str:
    return get_column_letter(n)


def frangobr_row(year: int, month: int) -> int:
    return 5 + (year - 2000) * 12 + (month - 1)


def production_row_to_ym(row: int):
    offset = row - 3
    return 2000 + offset // 12, offset % 12 + 1


def prev_month(year: int, month: int):
    return (year - 1, 12) if month == 1 else (year, month - 1)


def format_ym(year: int, month: int) -> str:
    return f"{month:02d}/{year}"


def last_formula_row(ws, col: int, start: int = 5) -> int:
    for r in range(ws.max_row, start - 1, -1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.startswith("="):
            return r
    return start - 1


def last_data_row(ws, col: int, start: int = 3) -> int:
    for r in range(ws.max_row, start - 1, -1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
            return r
    return 0


def last_nonempty_row(ws, col: int, start: int = 2) -> int:
    for r in range(ws.max_row, start - 1, -1):
        if ws.cell(row=r, column=col).value not in (None, ""):
            return r
    return start - 1


def extend_section(ws, src_row: int, col_ranges: list, dst_end: int) -> int:
    """Copia a linha-template (formula deslocada + formato) para baixo ate dst_end."""
    if dst_end <= src_row:
        return 0
    for c_start, c_end in col_ranges:
        for col in range(c_start, c_end + 1):
            src = ws.cell(row=src_row, column=col)
            src_val = src.value
            is_formula = isinstance(src_val, str) and src_val.startswith("=")
            for dst_row in range(src_row + 1, dst_end + 1):
                dst = ws.cell(row=dst_row, column=col)
                dst.value = shift_formula(src_val, dst_row - src_row) if is_formula else src_val
                dst.number_format = src.number_format
                dst.font = copy(src.font)
                dst.fill = copy(src.fill)
                dst.border = copy(src.border)
                dst.alignment = copy(src.alignment)
    return dst_end - src_row


# ── BBG_Dados ─────────────────────────────────────────────────────────────────

def atualizar_bbg_dados(ws_bbg) -> int:
    """Cols J, K, L alinhadas com a ultima linha com dado em col E."""
    last_data = last_nonempty_row(ws_bbg, col=5, start=2)
    last_jkl = last_formula_row(ws_bbg, col=10, start=2)

    if last_jkl < 2:
        print("  BBG_Dados J->L: nenhuma linha template em col J — pulando.")
        return 0

    if last_data > last_jkl:
        n = extend_section(ws_bbg, last_jkl, [(10, 12)], last_data)
        print(f"  BBG_Dados J->L: +{n} linha(s) -> linha {last_data}")
        return n
    print(f"  BBG_Dados J->L: ja atualizado (linha {last_jkl})")
    return 0


# ── FrangoBR ──────────────────────────────────────────────────────────────────

def atualizar_mes_corrente(ws_frango, today: datetime.date) -> int:
    """(i) Cols D->I, L, O->R, T — ate o mes corrente."""
    target = frangobr_row(today.year, today.month)
    last = last_formula_row(ws_frango, col=4)  # col D

    if last < 5:
        print("  (i)  Mes corrente: nenhuma linha template em col D — pulando.")
        return 0

    if target > last:
        n = extend_section(ws_frango, last, [(4, 9), (12, 12), (15, 18), (20, 20)], target)
        print(f"  (i)  Mes corrente: +{n} mes(es) -> {format_ym(today.year, today.month)}")
        return n
    print(f"  (i)  Mes corrente: ja atualizado (linha {last})")
    return 0


def atualizar_sif(ws_frango, ws_prod) -> int:
    """(ii) Col W — 2 meses antes do ultimo dado SIF em Production col F."""
    last_prod = last_data_row(ws_prod, col=6)
    if last_prod == 0:
        print("  (ii) SIF: sem dados em Production col F — pulando.")
        return 0

    yr, mo = production_row_to_ym(last_prod)
    yr, mo = prev_month(yr, mo)
    yr, mo = prev_month(yr, mo)
    target = frangobr_row(yr, mo)
    last = last_formula_row(ws_frango, col=23)  # col W

    if last < 5:
        print("  (ii) SIF: nenhuma linha template em col W — pulando.")
        return 0

    if target > last:
        n = extend_section(ws_frango, last, [(23, 23)], target)
        print(f"  (ii) SIF: +{n} mes(es) -> {format_ym(yr, mo)}")
        return n
    print(f"  (ii) SIF: ja atualizado (linha {last})")
    return 0


def atualizar_ibge(ws_frango, ws_prod) -> int:
    """(iii) Cols X, Y — ate o ultimo dado IBGE em Production col C."""
    last_prod = last_data_row(ws_prod, col=3)
    if last_prod == 0:
        print("  (iii) IBGE: sem dados em Production col C — pulando.")
        return 0

    target = last_prod + 2  # Production linha 3 = FrangoBR linha 5
    last = last_formula_row(ws_frango, col=24)  # col X

    if last < 5:
        print("  (iii) IBGE: nenhuma linha template em col X — pulando.")
        return 0

    if target > last:
        yr, mo = production_row_to_ym(last_prod)
        n = extend_section(ws_frango, last, [(24, 25)], target)
        print(f"  (iii) IBGE: +{n} mes(es) -> {format_ym(yr, mo)}")
        return n
    print(f"  (iii) IBGE: ja atualizado (linha {last})")
    return 0


def atualizar_apinco(ws_frango, ws_prod) -> int:
    """(iv) Col Z — ate o ultimo dado APINCO em Production col J."""
    last_prod = last_data_row(ws_prod, col=10)
    if last_prod == 0:
        print("  (iv) APINCO: sem dados em Production col J — pulando.")
        return 0

    target = last_prod + 2
    last = last_formula_row(ws_frango, col=26)  # col Z

    if last < 5:
        print("  (iv) APINCO: nenhuma linha template em col Z — pulando.")
        return 0

    if target > last:
        yr, mo = production_row_to_ym(last_prod)
        n = extend_section(ws_frango, last, [(26, 26)], target)
        print(f"  (iv) APINCO: +{n} mes(es) -> {format_ym(yr, mo)}")
        return n
    print(f"  (iv) APINCO: ja atualizado (linha {last})")
    return 0


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    fid = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(fid)

    today = datetime.date.today()
    print(f"FrangoBR — {today:%d/%m/%Y}\n")

    wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
    try:
        ws_bbg = wb["BBG_Dados"]
        ws_frango = wb["FrangoBR"]
        ws_prod = wb["Production"]

        total = 0
        total += atualizar_bbg_dados(ws_bbg)
        total += atualizar_mes_corrente(ws_frango, today)
        total += atualizar_sif(ws_frango, ws_prod)
        total += atualizar_ibge(ws_frango, ws_prod)
        total += atualizar_apinco(ws_frango, ws_prod)

        if total > 0:
            wb.save(path)
            drive_io.update_file(fid, path)   # sobe de volta pro Drive
            print(f"\nSalvo: {total} linha(s) estendida(s). "
                  f"(As formulas ganham valor quando o Excel/Bloomberg recalcular.)")
        else:
            print("\nNada a estender — ja atualizado.")
    finally:
        wb.close()


if __name__ == "__main__":
    main()
