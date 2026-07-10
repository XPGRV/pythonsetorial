import openpyxl
import pandas as pd
import datetime
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NIELSEN_PATH = r"C:\Users\gabri\Desktop\Nielsen_Dados_Mensais.xlsx"
DRIVE_ROOT = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
)
XLSX_NAME = "FrangoBR.xlsm"


def localizar_planilha() -> Path:
    """Localiza a planilha sem depender dos nomes das pastas intermediárias."""
    encontrados = list(DRIVE_ROOT.rglob(XLSX_NAME))

    if not encontrados:
        raise FileNotFoundError(
            f"Planilha {XLSX_NAME!r} não encontrada dentro de {DRIVE_ROOT}"
        )

    if len(encontrados) > 1:
        caminhos = "\n".join(f"  - {path}" for path in encontrados)
        raise RuntimeError(
            f"Foram encontradas várias planilhas chamadas {XLSX_NAME!r}:\n"
            f"{caminhos}"
        )

    return encontrados[0]


FRANGO_PATH = localizar_planilha()

# --- Colunas Nielsen na aba Processados ---
COL_START = 18   # R
NIELSEN_HEADERS = [
    "Industry - Vol. G%",
    "Industry - Px. G%",
    "BRF - Vol. G%",
    "BRF - Px. G%",
    "Seara - Vol. G%",
    "Seara - Px. G%",
]
NIELSEN_SECTION_TITLE = "Nielsen"

# --- Leitura Nielsen ---
df = pd.read_excel(NIELSEN_PATH, sheet_name="Dados Mensais")
df.columns = ["Mes", "Bimestre", "Ind_Vol", "Ind_Px", "BRF_Vol", "BRF_Px", "Seara_Vol", "Seara_Px"]

# Converte coluna Mes (mmm-yy) para datetime
month_map_pt = {
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12
}

def parse_mes(s):
    s = str(s).strip().lower()
    parts = s.split("-")
    if len(parts) != 2:
        return None
    mon = month_map_pt.get(parts[0])
    yr  = int(parts[1])
    yr  = yr + 2000 if yr < 100 else yr
    return datetime.datetime(yr, mon, 1) if mon else None

df["Data"] = df["Mes"].apply(parse_mes)
nielsen_map = {row["Data"]: row for _, row in df.iterrows()}

# --- Abertura FrangoBR ---
wb = openpyxl.load_workbook(FRANGO_PATH, keep_vba=True)
ws = wb["Processados"]

# --- Estilos replicados do cabeçalho existente ---
FILL_ROW1  = PatternFill("solid", fgColor="1E5880")   # azul escuro
FILL_ROW2  = PatternFill("solid", fgColor="3E97D2")   # azul médio
FONT_ROW1  = Font(bold=True, color="FFFFFF")
FONT_ROW2  = Font(bold=True, color="FFFFFF")
ALIGN_CTR  = Alignment(horizontal="center", vertical="center")

# --- Cabeçalho Linha 1: seção Nielsen (merged R1:W1) ---
col_end = COL_START + len(NIELSEN_HEADERS) - 1
merge_range = f"{get_column_letter(COL_START)}1:{get_column_letter(col_end)}1"

# Remove merge se já existir
for mc in list(ws.merged_cells.ranges):
    if str(mc) == merge_range:
        ws.unmerge_cells(merge_range)
        break

ws.merge_cells(merge_range)
cell_r1 = ws.cell(row=1, column=COL_START, value=NIELSEN_SECTION_TITLE)
cell_r1.font      = FONT_ROW1
cell_r1.fill      = FILL_ROW1
cell_r1.alignment = ALIGN_CTR

# Aplica fill nos demais cells da merge (openpyxl exige isso)
for c in range(COL_START + 1, col_end + 1):
    cll = ws.cell(row=1, column=c)
    cll.fill = FILL_ROW1

# --- Cabeçalho Linha 2: nomes das colunas ---
for i, hdr in enumerate(NIELSEN_HEADERS):
    cll = ws.cell(row=2, column=COL_START + i, value=hdr)
    cll.font      = FONT_ROW2
    cll.fill      = FILL_ROW2
    cll.alignment = ALIGN_CTR

# --- Preenchimento dos dados ---
updated = 0
for row in range(3, ws.max_row + 1):
    date_val = ws.cell(row=row, column=2).value
    if not isinstance(date_val, (datetime.datetime, datetime.date)):
        continue
    if isinstance(date_val, datetime.date) and not isinstance(date_val, datetime.datetime):
        date_val = datetime.datetime(date_val.year, date_val.month, 1)
    key = datetime.datetime(date_val.year, date_val.month, 1)
    if key not in nielsen_map:
        continue
    r = nielsen_map[key]
    values = [r["Ind_Vol"], r["Ind_Px"], r["BRF_Vol"], r["BRF_Px"], r["Seara_Vol"], r["Seara_Px"]]
    for i, v in enumerate(values):
        cll = ws.cell(row=row, column=COL_START + i)
        cll.value       = float(v)
        cll.number_format = "0.0%"
    updated += 1

wb.save(FRANGO_PATH)
print(f"Concluído. {updated} linhas atualizadas na aba Processados ({get_column_letter(COL_START)}-{get_column_letter(col_end)}).")
