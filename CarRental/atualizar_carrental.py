"""
atualizar_carrental.py  (versao openpyxl / nuvem)

Consolidador do CarRental — "puxa" as formulas-template para baixo, ate as
linhas novas.

Aba 'Preco Carros':
  Colunas F..M acompanham a ultima linha com %MoM coletado em D/E.
    F        = E - ajuste           (usado %MoM ajustado)
    G, H, I  = variacao, MM6M, MM12M
    J        = J[-1] * (1 + D)      (indice do novo, acumulado)
    K, L     = indices base 80
    M        = L / J - 1            (spread usado/novo)

Sem isto o coletor do SIDRA preenche D e E do mes novo e as colunas calculadas
ficam para tras. E aqui isso e' pior que no WEG: o dashboard le F, J, L e M, e a
validacao NAO pegaria -- ela conta pontos da serie, e as linhas historicas
continuam la, entao os minimos seguem satisfeitos enquanto o grafico envelhece.

Mapeamento de linhas: linha 5 = jan/2000 -> row = 5 + (ano-2000)*12 + (mes-1)

A aba 'Peers' nao e' tocada: as formulas de la sao BDH do Bloomberg, que se
espalham sozinhas quando o Excel recalcula.

Le/grava o CarRental.xlsm via Google Drive (drive_io). Roda na nuvem, sem Excel.
As formulas ganham valor quando o Excel (PC Bloomberg) recalcular.
Requer: pip install openpyxl google-api-python-client google-auth
"""

import re
import sys
from copy import copy
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import drive_io

XLSX_NAME = "CarRental.xlsm"
SHEET_NAME = "Preço Carros"
ROW_BASE = 5              # linha 5 = jan/2000

COL_NOVO = 4              # D — %MoM do automovel novo (coletado do SIDRA)
COL_USADO = 5             # E — %MoM do automovel usado (coletado do SIDRA)
COL_CALCULADAS = (6, 13)  # F..M — tudo que deriva de D e E

REF = re.compile(r"(?<![A-Za-z0-9_])(\$?[A-Za-z]{1,3})(\$?)(\d+)")


def shift(value, delta):
    """Desloca as referencias de linha de uma formula, respeitando os '$'."""
    if not isinstance(value, str) or not value.startswith("="):
        return value
    return REF.sub(
        lambda m: m.group(0)
        if m.group(2) == "$"
        else f"{m.group(1)}{m.group(2)}{int(m.group(3)) + delta}",
        value,
    )


def copy_cell(origem, destino, delta):
    destino.value = shift(origem.value, delta)
    destino._style = copy(origem._style)
    destino.number_format = origem.number_format
    destino.font = copy(origem.font)
    destino.fill = copy(origem.fill)
    destino.border = copy(origem.border)
    destino.alignment = copy(origem.alignment)


def last_formula(ws, col, start=ROW_BASE):
    """Ultima linha da coluna que ainda tem formula — serve de molde."""
    for row in range(ws.max_row, start - 1, -1):
        valor = ws.cell(row, col).value
        if isinstance(valor, str) and valor.startswith("="):
            return row
    return None


def last_numeric(ws, col, start=ROW_BASE):
    """Ultima linha com numero de verdade ('-' e vazio nao contam)."""
    ultima = 0
    for row in range(start, ws.max_row + 1):
        if isinstance(ws.cell(row, col).value, (int, float)):
            ultima = row
    return ultima


def extend(ws, src, end, ranges):
    if not src or end <= src:
        return 0
    changes = 0
    for row in range(src + 1, end + 1):
        for primeira, ultima in ranges:
            for col in range(primeira, ultima + 1):
                # So completa lacuna; nunca substitui o que ja esta escrito.
                if ws.cell(row, col).value not in (None, ""):
                    continue
                copy_cell(ws.cell(src, col), ws.cell(row, col), row - src)
                changes += 1
    return changes


def main():
    file_id = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(file_id)
    wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
    try:
        ws = wb[SHEET_NAME]

        # As calculadas vao ate onde houver %MoM coletado. Os dois meses chegam
        # juntos, mas o max evita depender disso.
        alvo = max(last_numeric(ws, COL_NOVO), last_numeric(ws, COL_USADO))
        molde = last_formula(ws, COL_CALCULADAS[0])

        changes = extend(ws, molde, alvo, (COL_CALCULADAS,))
        if changes:
            wb.save(path)
            drive_io.update_file(file_id, path)

        print(
            f"Consolidador CarRental: molde na linha {molde}, dados ate {alvo}, "
            f"{changes} celula(s) estendida(s)."
        )
    finally:
        wb.close()


if __name__ == "__main__":
    main()
