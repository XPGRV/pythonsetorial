"""
atualiza_production_frangous_rateolay.py

Baixa dados do USDA NASS QuickStats via API e atualiza a Coluna D
da aba "Production" no FrangoUS.xlsm:

  Coluna D  — Layers, Hatching, Broiler Type - Rate of Lay  (eggs/100 layers)

Fonte: USDA NASS QuickStats API  (https://quickstats.nass.usda.gov/api)
       Commodity : CHICKENS
       Short desc: CHICKENS, LAYERS, HATCHING, BROILER TYPE - RATE OF LAY
       Nível     : NATIONAL
"""

from __future__ import annotations

import datetime as dt
import os
import sys
import warnings
from pathlib import Path

import copy

import openpyxl
import requests

sys.stdout.reconfigure(encoding="utf-8")
warnings.filterwarnings("ignore")


# ── Configuração ──────────────────────────────────────────────────────────────

DRIVE_ROOT = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
)
XLSX_NAME = "FrangoUS.xlsm"
SHEET_NAME    = "Production"
FIRST_DATA_ROW = 3       # linha 3 = jan/2000
COL_DATE      = 2        # B — datetime mensal
COL_RATEOLAY  = 4        # D — Rate of Lay (eggs/100 layers)

# NASS_API_KEY vem do ambiente. Localmente, um .env (fora do git) ao lado
# deste script preenche o ambiente; na nuvem, use GitHub Secrets.
_env_file = Path(__file__).resolve().parent / ".env"
if _env_file.exists():
    for _line in _env_file.read_text(encoding="utf-8-sig").splitlines():
        _line = _line.strip()
        if _line and not _line.startswith("#") and "=" in _line:
            _k, _v = _line.split("=", 1)
            os.environ.setdefault(_k.strip(), _v.strip().strip('"').strip("'"))

NASS_API_KEY = os.environ.get("NASS_API_KEY")
if not NASS_API_KEY:
    raise SystemExit("Defina NASS_API_KEY (variavel de ambiente ou .env).")

NASS_BASE_URL = "https://quickstats.nass.usda.gov/api/api_GET/"

NASS_PARAMS = {
    "source_desc":      "SURVEY",
    "sector_desc":      "ANIMALS & PRODUCTS",
    "commodity_desc":   "CHICKENS",
    "short_desc":       "CHICKENS, LAYERS, HATCHING, BROILER TYPE - RATE OF LAY, MEASURED IN EGGS / 100 LAYER",
    "agg_level_desc":   "NATIONAL",
    "freq_desc":        "MONTHLY",
    "format":           "JSON",
}

MONTH_MAP = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12,
}


# ── 1. Download NASS ──────────────────────────────────────────────────────────

def localizar_planilha() -> Path:
    """Localiza a planilha sem depender dos nomes das pastas intermediárias."""
    encontrados = list(DRIVE_ROOT.rglob(XLSX_NAME))

    if not encontrados:
        raise FileNotFoundError(
            f"Planilha {XLSX_NAME!r} não encontrada dentro de {DRIVE_ROOT}"
        )

    if len(encontrados) > 1:
        caminhos = "\n".join(f"  - {path}" for path in encontrados)
        raise RuntimeError(
            f"Foram encontradas várias planilhas chamadas {XLSX_NAME!r}:\n"
            f"{caminhos}"
        )

    return encontrados[0]


def fetch_nass() -> list[dict]:
    """Chama a API NASS e retorna a lista de registros JSON."""
    params = {**NASS_PARAMS, "key": NASS_API_KEY}
    print(f"Consultando NASS QuickStats: {NASS_PARAMS['short_desc']}")
    resp = requests.get(NASS_BASE_URL, params=params, timeout=60)

    if resp.status_code == 200:
        data = resp.json()
        if "data" in data:
            print(f"  {len(data['data'])} registros recebidos.")
            return data["data"]
        raise RuntimeError(f"Resposta NASS inesperada: {data}")

    # HTTP 413 = muitos registros — tenta com filtro de ano mínimo
    if resp.status_code == 413:
        print("  Muitos registros — adicionando filtro year__GE=2000 ...")
        params["year__GE"] = "2000"
        resp = requests.get(NASS_BASE_URL, params=params, timeout=60)
        resp.raise_for_status()
        data = resp.json()
        if "data" in data:
            print(f"  {len(data['data'])} registros recebidos.")
            return data["data"]

    resp.raise_for_status()
    raise RuntimeError("Resposta NASS inesperada.")


# ── 2. Parse ──────────────────────────────────────────────────────────────────

def parse_nass(records: list[dict]) -> dict[tuple[int, int], float]:
    """
    Converte lista de registros NASS em {(ano, mes): valor_em_eggs_por_100_layers}.

    A unidade retornada pela API é "EGGS / 100 LAYERS" — sem conversão.
    Aceita reference_period_desc tanto no formato "FIRST OF JAN" quanto "JAN".
    """
    result: dict[tuple[int, int], float] = {}
    skipped = 0

    for rec in records:
        period = str(rec.get("reference_period_desc", "")).strip().upper()
        year_s = str(rec.get("year", "")).strip()
        value_s = str(rec.get("Value", "")).strip().replace(",", "")

        # Período pode vir como 'JAN', 'FEB' ou 'FIRST OF JAN', etc.
        month = MONTH_MAP.get(period[-3:])
        if month is None:
            skipped += 1
            continue

        try:
            year = int(year_s)
            value = float(value_s)
        except (ValueError, TypeError):
            skipped += 1
            continue

        result[(year, month)] = value

    print(f"  Parse: {len(result)} meses mensais válidos ({skipped} ignorados).")
    if result:
        min_ym = min(result)
        max_ym = max(result)
        print(f"  Intervalo: {min_ym[0]}-{min_ym[1]:02d}  →  {max_ym[0]}-{max_ym[1]:02d}")
    return result


# ── 3. Atualização da planilha ────────────────────────────────────────────────

def _val_is_empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and v.strip().lower() in ("", "sem dados"):
        return True
    return False


def _approx_equal(a: float, b: float, tol: float = 1e-9) -> bool:
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def update_production(nass_data: dict[tuple[int, int], float]) -> dict:
    path = localizar_planilha()

    wb = openpyxl.load_workbook(str(path), keep_vba=True, data_only=False)
    try:
        ws = wb[SHEET_NAME]

        # Referência de formatação: última célula preenchida com formato correto.
        ref_fmt = None
        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            cell = ws.cell(row=row, column=COL_RATEOLAY)
            if cell.value is not None and cell.number_format not in ("General", ""):
                ref_fmt = cell
        if ref_fmt is None:
            for row in range(FIRST_DATA_ROW, ws.max_row + 1):
                cell = ws.cell(row=row, column=COL_RATEOLAY)
                if cell.number_format not in ("General", ""):
                    ref_fmt = cell
                    break

        def _apply_fmt(cell) -> None:
            if ref_fmt is None:
                return
            cell.font          = copy.copy(ref_fmt.font)
            cell.fill          = copy.copy(ref_fmt.fill)
            cell.border        = copy.copy(ref_fmt.border)
            cell.alignment     = copy.copy(ref_fmt.alignment)
            cell.number_format = ref_fmt.number_format

        novos    = 0
        revisoes = 0
        fmt_fix  = 0
        updates: list[str] = []

        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            date_val = ws.cell(row=row, column=COL_DATE).value
            if date_val is None:
                continue
            if not isinstance(date_val, (dt.datetime, dt.date)):
                continue

            ym = (date_val.year, date_val.month)
            new_val = nass_data.get(ym)
            if new_val is None:
                continue

            cell = ws.cell(row=row, column=COL_RATEOLAY)
            old_val = cell.value

            if _val_is_empty(old_val):
                cell.value = new_val
                _apply_fmt(cell)
                novos += 1
                updates.append(f"  [{ym[0]}-{ym[1]:02d}] D: (vazio) → {new_val:.2f}")
            elif isinstance(old_val, (int, float)):
                if not _approx_equal(float(old_val), new_val):
                    cell.value = new_val
                    _apply_fmt(cell)
                    revisoes += 1
                    updates.append(
                        f"  [{ym[0]}-{ym[1]:02d}] D: {float(old_val):.2f} → {new_val:.2f}  (revisão)"
                    )
                elif cell.number_format in ("General", ""):
                    _apply_fmt(cell)
                    fmt_fix += 1

        total = novos + revisoes + fmt_fix
        if total > 0:
            wb.save(str(path))
            msg_parts = []
            if novos:    msg_parts.append(f"{novos} nova(s)")
            if revisoes: msg_parts.append(f"{revisoes} revisão/revisões")
            if fmt_fix:  msg_parts.append(f"{fmt_fix} formato(s) corrigido(s)")
            print(f"Salvo. {total} célula(s) alterada(s) ({', '.join(msg_parts)}).")
            for u in updates[:60]:
                print(u)
            if len(updates) > 60:
                print(f"  ... (+{len(updates) - 60} mais)")
        else:
            print("Nenhuma alteração necessária — planilha já está atualizada.")

        return {"alteracoes": total, "novos": novos, "revisoes": revisoes, "fmt_fix": fmt_fix}
    finally:
        wb.close()


# ── 4. Main ───────────────────────────────────────────────────────────────────

def main() -> None:
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] atualiza_production_frangous_rateolay.py")
    print(f"Arquivo : {XLSX_NAME}")
    print(f"Aba     : {SHEET_NAME}  |  Coluna D — Rate of Lay (eggs/100 layers)")

    records   = fetch_nass()
    nass_data = parse_nass(records)

    if not nass_data:
        print("Nenhum dado mensal encontrado. Verifique os parâmetros NASS.")
        return

    print(f"\nAtualizando {XLSX_NAME} ...")
    result = update_production(nass_data)

    print(
        f"\n[{dt.datetime.now():%Y-%m-%d %H:%M}] Concluído. "
        f"{result['alteracoes']} alteração(ões) "
        f"({result['novos']} nova(s), {result['revisoes']} revisão/revisões)."
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit(130)
