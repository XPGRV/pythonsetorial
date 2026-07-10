"""
atualiza_wholesale_broiler_composite.py

Baixa os dados do site USDA ERS (Meat Price Spreads) e atualiza a coluna T
("Wholesale Broiler Composite", USDc/lb) da aba FrangoUS no FrangoUS.xlsm.

Duas fontes:
  - Histórico : "Historical monthly price spread data for beef, pork, broilers"
                aba "broiler", coluna B (Wholesale broiler composite)
                Cobertura: jan/1990 em diante (mais longa, maior precisão)

  - Recente   : "Retail prices for beef, pork, poultry cuts, eggs, and dairy products"
                aba "cuts", coluna 21 (0-indexed) (Wholesale broiler composite)
                Cobertura: mar/2024 em diante (dados mais recentes)

Regra de prioridade: o Histórico sempre prevalece. O Recente preenche apenas
os meses ainda não cobertos pelo Histórico.

Requer: pip install requests openpyxl
"""

from __future__ import annotations

import datetime as dt
import io
import sys
import warnings
from pathlib import Path

import openpyxl
import requests

sys.stdout.reconfigure(encoding="utf-8")
warnings.filterwarnings("ignore")


# ── Caminhos ──────────────────────────────────────────────────────────────────
DRIVE_ROOT = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
)
XLSX_NAME = "FrangoUS.xlsm"

# ── Planilha FrangoUS ─────────────────────────────────────────────────────────
SHEET_NAME     = "FrangoUS"
FIRST_DATA_ROW = 5      # primeira linha com dados (logo abaixo dos cabeçalhos)
COL_DATE       = 2      # B — datetime mensal
COL_T          = 20     # T — Wholesale Broiler Composite (USDc/lb)

# ── URLs USDA ERS ─────────────────────────────────────────────────────────────
URL_HISTORICAL = (
    "https://www.ers.usda.gov/media/5027/"
    "historical-monthly-price-spread-data-for-beef-pork-broilers.xlsx"
)
URL_RETAIL = (
    "https://www.ers.usda.gov/media/5021/"
    "retail-prices-for-beef-pork-poultry-cuts-eggs-and-dairy-products.xlsx"
)


# ─────────────────────────────────────────────────────────────────────────────
# 1. Download
# ─────────────────────────────────────────────────────────────────────────────

def localizar_planilha() -> Path:
    """Localiza a planilha sem depender dos nomes das pastas intermediárias."""
    encontrados = list(DRIVE_ROOT.rglob(XLSX_NAME))

    if not encontrados:
        raise FileNotFoundError(
            f"Planilha {XLSX_NAME!r} não encontrada dentro de {DRIVE_ROOT}"
        )

    if len(encontrados) > 1:
        caminhos = "\n".join(f"  - {path}" for path in encontrados)
        raise RuntimeError(
            f"Foram encontradas várias planilhas chamadas {XLSX_NAME!r}:\n"
            f"{caminhos}"
        )

    return encontrados[0]


def _session() -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "User-Agent": "Mozilla/5.0 (compatible; FrangoUS-updater/1.0)",
        "Accept-Language": "en-US,en;q=0.9",
    })
    return s


def download(label: str, url: str) -> bytes:
    s = _session()
    print(f"Baixando {label}: {url}")
    resp = s.get(url, timeout=120)
    resp.raise_for_status()
    print(f"  {len(resp.content) / 1024:.1f} KB recebidos.")
    return resp.content


# ─────────────────────────────────────────────────────────────────────────────
# 2. Parse
# ─────────────────────────────────────────────────────────────────────────────

def _to_float(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if f != f else f   # NaN check
    try:
        return float(str(v).strip())
    except (ValueError, TypeError):
        return None


def _to_ym(v) -> tuple[int, int] | None:
    """Extrai (year, month) de um valor de data, normalizando para o 1º do mês."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year, v.month
    return None


def parse_historical(xlsx_bytes: bytes) -> dict[tuple[int, int], float]:
    """
    Lê a aba 'broiler' do arquivo histórico.
    Retorna {(year, month): wholesale_value}.
    Os 4 primeiros rows são cabeçalho; col index 1 (B) = Wholesale.
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb["broiler"]

    result: dict[tuple[int, int], float] = {}
    header_skipped = 0

    for row in ws.iter_rows(min_row=1, values_only=True):
        # Pula as 4 linhas de cabeçalho (título, unidade, labels, "Date")
        if header_skipped < 4:
            header_skipped += 1
            continue

        date_val  = row[0]   # coluna A = Date
        wholesale = row[1]   # coluna B = Wholesale broiler composite

        ym = _to_ym(date_val)
        if ym is None:
            continue

        val = _to_float(wholesale)
        if val is not None:
            result[ym] = val

    wb.close()
    print(f"  Histórico: {len(result)} meses com Wholesale Broiler Composite.")
    if result:
        print(f"  Cobertura: {min(result)} → {max(result)}")
    return result


def parse_retail(xlsx_bytes: bytes) -> dict[tuple[int, int], float]:
    """
    Lê a aba 'cuts' do arquivo de preços ao varejo.
    Retorna {(year, month): wholesale_value}.
    Os 5 primeiros rows são cabeçalho; col index 21 (0-based) = Wholesale broiler composite.
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb["cuts"]

    result: dict[tuple[int, int], float] = {}
    header_skipped = 0

    for row in ws.iter_rows(min_row=1, values_only=True):
        # Pula as 5 linhas de cabeçalho
        if header_skipped < 5:
            header_skipped += 1
            continue

        date_val  = row[0]    # coluna A = date
        wholesale = row[21]   # coluna V (0-indexed 21) = Wholesale broiler composite

        ym = _to_ym(date_val)
        if ym is None:
            continue

        val = _to_float(wholesale)
        if val is not None:
            result[ym] = val

    wb.close()
    print(f"  Recente: {len(result)} meses com Wholesale Broiler Composite.")
    if result:
        print(f"  Cobertura: {min(result)} → {max(result)}")
    return result


def merge(
    historical: dict[tuple[int, int], float],
    retail: dict[tuple[int, int], float],
) -> dict[tuple[int, int], float]:
    """
    Mescla as duas fontes com prioridade para o histórico.
    Começa com os dados recentes e sobrescreve com o histórico.
    """
    merged = {**retail, **historical}   # histórico prevalece
    only_retail = len(set(retail) - set(historical))
    print(
        f"  Mesclado: {len(merged)} meses no total "
        f"({len(historical)} do histórico, {only_retail} apenas do recente)."
    )
    return merged


# ─────────────────────────────────────────────────────────────────────────────
# 3. Atualização do FrangoUS.xlsm
# ─────────────────────────────────────────────────────────────────────────────

def _val_is_empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and v.strip().lower() in ("sem dados", ""):
        return True
    return False


def _approx_equal(a: float, b: float, tol: float = 1e-9) -> bool:
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def update_frangous(data: dict[tuple[int, int], float]) -> dict:
    """
    Abre FrangoUS.xlsm (keep_vba=True) e atualiza a coluna T.
    """
    # Baixa a planilha do Google Drive (em vez do drive montado em G:\).
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import drive_io
    fid = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(fid)

    wb = openpyxl.load_workbook(str(path), keep_vba=True, data_only=False)
    try:
        ws = wb[SHEET_NAME]

        updates: list[str] = []
        novos    = 0
        revisoes = 0

        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            date_val = ws.cell(row=row, column=COL_DATE).value
            if date_val is None:
                continue
            if not isinstance(date_val, (dt.datetime, dt.date)):
                continue

            ym = (date_val.year, date_val.month)
            new_val = data.get(ym)
            if new_val is None:
                continue

            cell    = ws.cell(row=row, column=COL_T)
            old_val = cell.value

            if _val_is_empty(old_val):
                cell.value = new_val
                novos += 1
                updates.append(
                    f"  [{ym[0]}-{ym[1]:02d}] T: '{old_val}' -> {new_val}"
                )
            elif isinstance(old_val, (int, float)):
                if not _approx_equal(float(old_val), new_val):
                    cell.value = new_val
                    revisoes += 1
                    updates.append(
                        f"  [{ym[0]}-{ym[1]:02d}] T: {float(old_val)} -> {new_val} (revisão)"
                    )

        total = novos + revisoes
        if total > 0:
            wb.save(str(path)); drive_io.update_file(fid, path)
            print(f"Salvo. {total} célula(s) alterada(s) ({novos} nova(s), {revisoes} revisão/revisões).")
            for u in updates[:50]:
                print(u)
            if len(updates) > 50:
                print(f"  ... (+{len(updates) - 50} mais)")
        else:
            print("Nenhuma alteração necessária — planilha já está atualizada.")

        return {"alteracoes": total, "novos": novos, "revisoes": revisoes}
    finally:
        wb.close()


# ─────────────────────────────────────────────────────────────────────────────
# 4. Main
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] atualiza_wholesale_broiler_composite.py")
    print(f"Arquivo: {XLSX_NAME}\n")

    hist_bytes   = download("Histórico", URL_HISTORICAL)
    retail_bytes = download("Recente",   URL_RETAIL)

    print("\nParseando dados...")
    historical = parse_historical(hist_bytes)
    retail     = parse_retail(retail_bytes)

    print("\nMesclando fontes (histórico tem prioridade)...")
    data = merge(historical, retail)

    print(f"\nAtualizando {XLSX_NAME} — coluna T (Wholesale Broiler Composite)...")
    result = update_frangous(data)

    print(
        f"\n[{dt.datetime.now():%Y-%m-%d %H:%M}] Concluído. "
        f"{result['alteracoes']} alteração(ões) "
        f"({result['novos']} nova(s), {result['revisoes']} revisão/revisões)."
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit(130)
