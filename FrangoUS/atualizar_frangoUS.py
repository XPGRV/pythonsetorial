"""
atualizar_frangoUS.py  (versao openpyxl / nuvem)

Consolidador do FrangoUS — "puxa" as formulas-template para baixo, ate as
linhas novas, em duas abas:

Aba BBG_Dados:
  Colunas I, J, L, O, Q, R, V, X, AA — ate a ultima linha com dado em col E.

Aba FrangoUS:
  (i)   Mes corrente:  D->F, I->P, W, AP, AQ  — ate o mes corrente
  (ii)  Mes anterior:  AE->AH, AK->AN, AR      — sempre o mes anterior ao corrente
  (iii) Prod col C:    Z                        — ultimo dado em Production col C
  (iv)  Prod col D:    AA->AD                   — ultimo dado em Production col D
  (v)   Prod col G:    AI, AJ                   — ultimo dado em Production col G
  (vi)  Prod col H:    AO                       — ultimo dado em Production col H

Mapeamento de linhas:
  FrangoUS:   linha 5 = jan/2000  ->  row = 5 + (ano-2000)*12 + (mes-1)
  Production: linha 3 = jan/2000  ->  FrangoUS_row = Production_row + 2

As formulas estendidas NAO sao BDH — sao aritmeticas referenciando outras
colunas, todas relativas (sem $). "Puxar pra baixo" = incrementar o numero de
linha de cada referencia (o mesmo que o Excel faz num arrasta/cola). As celulas
novas ficam com formula sem valor em cache ate o Excel (PC Bloomberg) recalcular.

Le/grava o FrangoUS.xlsm via Google Drive (drive_io). Roda na nuvem, sem Excel.
Requer: pip install openpyxl google-api-python-client google-auth
"""

import re
import sys
import datetime
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

XLSX_NAME = "FrangoUS.xlsm"


# ── Deslocamento de formula ("puxar pra baixo") ───────────────────────────────
# Incrementa o numero de linha de cada referencia relativa. Referencias com $
# na linha (absolutas) nao mudam. Validado celula a celula contra o Excel real.
_REF = re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?![A-Za-z_(])")


def shift_formula(formula, delta: int):
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(m):
        col_abs, col, row_abs, row = m.groups()
        if row_abs == "$":
            return m.group(0)
        return f"{col_abs}{col}{row_abs}{int(row) + delta}"

    return _REF.sub(repl, formula)


# ── Helpers de linha ──────────────────────────────────────────────────────────

def col_letter(n: int) -> str:
    return get_column_letter(n)


def frangoUS_row(year: int, month: int) -> int:
    """Linha no FrangoUS para ano/mes. Linha 5 = jan/2000."""
    return 5 + (year - 2000) * 12 + (month - 1)


def production_row_to_ym(row: int):
    """Linha da aba Production -> (ano, mes). Linha 3 = jan/2000."""
    offset = row - 3
    return 2000 + offset // 12, offset % 12 + 1


def prev_month(year: int, month: int):
    return (year - 1, 12) if month == 1 else (year, month - 1)


def format_ym(year: int, month: int) -> str:
    return f"{month:02d}/{year}"


def last_formula_row(ws, col: int, start: int = 5) -> int:
    """Ultima linha em 'col' com formula, de baixo para cima."""
    for r in range(ws.max_row, start - 1, -1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.startswith("="):
            return r
    return start - 1


def last_data_row(ws, col: int, start: int = 3) -> int:
    """Ultima linha em 'col' com valor numerico > 0, de baixo para cima."""
    for r in range(ws.max_row, start - 1, -1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
            return r
    return 0


def last_nonempty_row(ws, col: int, start: int = 2) -> int:
    """Ultima linha em 'col' com qualquer conteudo, de baixo para cima."""
    for r in range(ws.max_row, start - 1, -1):
        if ws.cell(row=r, column=col).value not in (None, ""):
            return r
    return start - 1


def extend_section(ws, src_row: int, col_ranges: list, dst_end: int) -> int:
    """Copia a linha-template 'src_row' (formula deslocada + formato) para baixo
    ate 'dst_end', nas colunas dadas. Equivale ao Copy + PasteSpecial(All) do Excel."""
    if dst_end <= src_row:
        return 0
    for c_start, c_end in col_ranges:
        for col in range(c_start, c_end + 1):
            src = ws.cell(row=src_row, column=col)
            src_val = src.value
            is_formula = isinstance(src_val, str) and src_val.startswith("=")
            for dst_row in range(src_row + 1, dst_end + 1):
                dst = ws.cell(row=dst_row, column=col)
                dst.value = shift_formula(src_val, dst_row - src_row) if is_formula else src_val
                dst.number_format = src.number_format
                dst.font = copy(src.font)
                dst.fill = copy(src.fill)
                dst.border = copy(src.border)
                dst.alignment = copy(src.alignment)
    return dst_end - src_row


# ── BBG_Dados ─────────────────────────────────────────────────────────────────

def atualizar_bbg_dados(ws_bbg) -> int:
    """Cols I, J, L, O, Q, R, V, X, AA — ultima linha com dado em col E."""
    last_data = last_nonempty_row(ws_bbg, col=5, start=2)
    last_tmpl = last_formula_row(ws_bbg, col=9, start=2)  # col I

    if last_tmpl < 2:
        print("  BBG_Dados: nenhuma linha template em col I — pulando.")
        return 0

    if last_data > last_tmpl:
        n = extend_section(
            ws_bbg, last_tmpl,
            [(9, 10), (12, 12), (15, 15), (17, 18), (22, 22), (24, 24), (27, 27)],
            last_data,
        )
        print(f"  BBG_Dados: +{n} linha(s) -> linha {last_data}")
        return n
    print(f"  BBG_Dados: ja atualizado (linha {last_tmpl})")
    return 0


# ── FrangoUS ──────────────────────────────────────────────────────────────────

def atualizar_mes_corrente(ws_us, today: datetime.date) -> int:
    """(i) Cols D->F, I->P, W, AP, AQ — ate o mes corrente."""
    target = frangoUS_row(today.year, today.month)
    last   = last_formula_row(ws_us, col=4)  # col D

    if last < 5:
        print("  (i)  Mes corrente: nenhuma linha template em col D — pulando.")
        return 0

    if target > last:
        n = extend_section(ws_us, last, [(4, 6), (9, 16), (23, 23), (42, 43)], target)
        print(f"  (i)  Mes corrente: +{n} mes(es) -> {format_ym(today.year, today.month)}")
        return n
    print(f"  (i)  Mes corrente: ja atualizado (linha {last})")
    return 0


def atualizar_mes_anterior(ws_us, today: datetime.date) -> int:
    """(ii) Cols AE->AH, AK->AN, AR — sempre o mes anterior ao corrente."""
    yr, mo = prev_month(today.year, today.month)
    target = frangoUS_row(yr, mo)
    last   = last_formula_row(ws_us, col=31)  # col AE

    if last < 5:
        print("  (ii) Mes anterior: nenhuma linha template em col AE — pulando.")
        return 0

    if target > last:
        n = extend_section(ws_us, last, [(31, 34), (37, 40), (44, 44)], target)
        print(f"  (ii) Mes anterior: +{n} mes(es) -> {format_ym(yr, mo)}")
        return n
    print(f"  (ii) Mes anterior: ja atualizado (linha {last})")
    return 0


def atualizar_prod_col(ws_us, ws_prod, prod_col: int, frango_col_ranges: list,
                       label: str) -> int:
    """Generico: estende colunas de ws_us ate o ultimo dado em ws_prod[prod_col]."""
    last_prod = last_data_row(ws_prod, col=prod_col)

    if last_prod == 0:
        print(f"  {label}: sem dados em Production col {col_letter(prod_col)} — pulando.")
        return 0

    target    = last_prod + 2  # Production linha 3 = FrangoUS linha 5
    ref_col   = frango_col_ranges[0][0]
    last_tmpl = last_formula_row(ws_us, col=ref_col)

    if last_tmpl < 5:
        print(f"  {label}: nenhuma linha template em col {col_letter(ref_col)} — pulando.")
        return 0

    if target > last_tmpl:
        yr, mo = production_row_to_ym(last_prod)
        n = extend_section(ws_us, last_tmpl, frango_col_ranges, target)
        print(f"  {label}: +{n} mes(es) -> {format_ym(yr, mo)}")
        return n
    print(f"  {label}: ja atualizado (linha {last_tmpl})")
    return 0


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    # Baixa a planilha do Google Drive (em vez do drive montado em G:\).
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import drive_io
    fid  = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(fid)

    today = datetime.date.today()
    print(f"FrangoUS — {today:%d/%m/%Y}\n")

    wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
    try:
        ws_bbg  = wb["BBG_Dados"]
        ws_us   = wb["FrangoUS"]
        ws_prod = wb["Production"]

        total = 0
        total += atualizar_bbg_dados(ws_bbg)
        total += atualizar_mes_corrente(ws_us, today)
        total += atualizar_mes_anterior(ws_us, today)
        total += atualizar_prod_col(ws_us, ws_prod, 3, [(26, 26)], "(iii) Prod C -> Z")
        total += atualizar_prod_col(ws_us, ws_prod, 4, [(27, 30)], "(iv)  Prod D -> AA->AD")
        total += atualizar_prod_col(ws_us, ws_prod, 7, [(35, 36)], "(v)   Prod G -> AI,AJ")
        total += atualizar_prod_col(ws_us, ws_prod, 8, [(41, 41)], "(vi)  Prod H -> AO")

        if total > 0:
            wb.save(path)
            drive_io.update_file(fid, path)   # sobe de volta pro Drive (in-place)
            print(f"\nSalvo: {total} linha(s) estendida(s). "
                  f"(As formulas ganham valor quando o Excel/Bloomberg recalcular.)")
        else:
            print("\nNada a estender — ja atualizado.")
    finally:
        wb.close()


if __name__ == "__main__":
    main()
