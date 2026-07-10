"""
atualiza_usda_frangous.py

Baixa o arquivo "Production Indicators and Estimated Returns" do USDA ERS e
atualiza as colunas Q, R, S da aba FrangoUS no FrangoUS.xlsm com os indices
de broiler (base 1998-2000=100):

  Coluna Q (17): Broilers, Feed Costs per Lb (Live Weight)
  Coluna R (18): Broilers, Composite Wholesale Price
  Coluna S (19): Broilers, Composite Price minus Feed Costs

O script verifica data por data e so escreve quando o valor do USDA for
diferente do que ja esta na planilha (captura tanto dados novos quanto revisoes).

Requer: pip install requests openpyxl
"""

from __future__ import annotations

import datetime as dt
import io
import re
import sys
import warnings
from pathlib import Path

import openpyxl
import requests

sys.stdout.reconfigure(encoding="utf-8")
warnings.filterwarnings("ignore")


# ── Caminhos ─────────────────────────────────────────────────────────────────
DRIVE_ROOT = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
)
XLSX_NAME = "FrangoUS.xlsm"

# ── Planilha FrangoUS ────────────────────────────────────────────────────────
SHEET_NAME      = "FrangoUS"
FIRST_DATA_ROW  = 5     # primeira linha com dados (logo abaixo dos cabecalhos)
COL_DATE        = 2     # B — datetime mensal
COL_FEED_COST   = 17    # Q
COL_WHOLESALE   = 18    # R
COL_SPREAD      = 19    # S

# ── USDA ERS ──────────────────────────────────────────────────────────────────
ERS_PAGE_URL    = "https://www.ers.usda.gov/data-products/livestock-and-meat-domestic-data"
DIRECT_XLSX_URL = "https://www.ers.usda.gov/media/20292/production-indicators-and-estimated-returns.xlsx"

# Rotulos buscados dentro do bloco "Broiler" no arquivo USDA
LABEL_SECTION   = "broiler (index"   # especifico para nao confundir com "Broilers (1,000)"
LABEL_FEED      = "feed cost"
LABEL_MARKET    = "market price"
LABEL_PRICE_LESS = "price less cost"


# ─────────────────────────────────────────────────────────────────────────────
# 1. Download do arquivo USDA
# ─────────────────────────────────────────────────────────────────────────────

def localizar_planilha() -> Path:
    """Localiza a planilha sem depender dos nomes das pastas intermediarias."""
    encontrados = list(DRIVE_ROOT.rglob(XLSX_NAME))

    if not encontrados:
        raise FileNotFoundError(
            f"Planilha {XLSX_NAME!r} nao encontrada dentro de {DRIVE_ROOT}"
        )

    if len(encontrados) > 1:
        caminhos = "\n".join(f"  - {path}" for path in encontrados)
        raise RuntimeError(
            f"Foram encontradas varias planilhas chamadas {XLSX_NAME!r}:\n"
            f"{caminhos}"
        )

    return encontrados[0]


def _session() -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "User-Agent": "Mozilla/5.0 (compatible; FrangoUS-updater/1.0)",
        "Accept-Language": "en-US,en;q=0.9",
    })
    return s


def _find_xlsx_url(page_url: str) -> str | None:
    """Tenta encontrar o link do XLSX na pagina do ERS."""
    try:
        s = _session()
        html = s.get(page_url, timeout=30).text
        m = re.search(
            r'href="([^"]*production-indicators[^"]*\.xlsx[^"]*)"',
            html,
            flags=re.IGNORECASE,
        )
        if m:
            href = m.group(1)
            if href.startswith("http"):
                return href
            return "https://www.ers.usda.gov" + href
    except Exception as e:
        print(f"[WARN] Falha ao buscar URL na pagina ERS: {e}")
    return None


def download_usda_xlsx() -> bytes:
    """Baixa o XLSX do USDA ERS. Tenta primeiro scrapar o URL da pagina."""
    s = _session()
    url = _find_xlsx_url(ERS_PAGE_URL) or DIRECT_XLSX_URL
    print(f"Baixando USDA ERS: {url}")
    resp = s.get(url, timeout=120)
    resp.raise_for_status()
    print(f"  {len(resp.content) / 1024:.1f} KB recebidos.")
    return resp.content


# ─────────────────────────────────────────────────────────────────────────────
# 2. Parse do arquivo USDA
# ─────────────────────────────────────────────────────────────────────────────

def _cell_to_float(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if f != f else f   # NaN check
    try:
        return float(str(v).strip())
    except (ValueError, TypeError):
        return None


def _cell_to_ym(v) -> tuple[int, int] | None:
    """Extrai (year, month) de um valor de cabecalho de data do USDA."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year, v.month
    if isinstance(v, str):
        v = v.strip()
        months_en = {
            "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
            "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
            "january": 1, "february": 2, "march": 3, "april": 4,
            "june": 6, "july": 7, "august": 8, "september": 9,
            "october": 10, "november": 11, "december": 12,
        }
        m = re.match(
            r"([A-Za-z]+)[.\-\s]+(\d{4})|(\d{4})[.\-\s]+([A-Za-z]+)",
            v,
        )
        if m:
            mon_s = (m.group(1) or m.group(4)).lower()
            yr_s  = m.group(2) or m.group(3)
            mon   = months_en.get(mon_s[:3])
            if mon:
                return int(yr_s), mon
    return None


def parse_usda_broiler_data(xlsx_bytes: bytes) -> dict[tuple[int, int], tuple[float, float, float]]:
    """
    Le o XLSX do USDA ERS e retorna:
      {(year, month): (feed_cost, market_price, price_less_cost)}
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb.active
    print(f"  Aba USDA: '{ws.title}'  ({ws.max_row} linhas x {ws.max_column} colunas)")

    max_r = ws.max_row
    max_c = ws.max_column

    # Encontra a linha de cabecalho da secao "Estimated returns" (datas do bloco broiler).
    # O arquivo tem dois cabecalhos de data: linha 1 (Production indicators) e uma linha
    # dentro da secao "Estimated returns" — que e a correta para Feed cost/Market price/Spread.
    estimated_returns_row: int | None = None
    for r in range(1, max_r + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and "estimated returns" in str(v).strip().lower():
            estimated_returns_row = r
            break

    # Header de datas: a propria linha "Estimated returns" ja contem as datas nas colunas B+
    header_row: int | None = None
    if estimated_returns_row is not None:
        count = sum(
            1 for c in range(2, min(max_c, 10) + 1)
            if _cell_to_ym(ws.cell(row=estimated_returns_row, column=c).value) is not None
        )
        if count >= 3:
            header_row = estimated_returns_row

    # Fallback: primeira linha com datas (caso o layout mude)
    if header_row is None:
        for r in range(1, min(10, max_r) + 1):
            count = sum(
                1 for c in range(2, min(max_c, 10) + 1)
                if _cell_to_ym(ws.cell(row=r, column=c).value) is not None
            )
            if count >= 3:
                header_row = r
                break

    if header_row is None:
        raise RuntimeError("Nao encontrei a linha de datas no arquivo USDA.")

    print(f"  Header de datas: linha {header_row}")

    # Mapeia coluna -> (year, month)
    col_to_ym: dict[int, tuple[int, int]] = {}
    for c in range(2, max_c + 1):
        ym = _cell_to_ym(ws.cell(row=header_row, column=c).value)
        if ym is not None:
            col_to_ym[c] = ym

    min_ym = min(col_to_ym.values())
    max_ym = max(col_to_ym.values())
    print(f"  Datas mapeadas: {len(col_to_ym)} meses ({min_ym} -> {max_ym})")

    # Encontra as linhas Feed cost / Market price / Price less cost do bloco Broiler
    row_feed: int | None = None
    row_market: int | None = None
    row_spread: int | None = None
    in_broiler = False

    for r in range(header_row + 1, max_r + 1):
        raw = ws.cell(row=r, column=1).value
        if raw is None:
            continue
        label = str(raw).strip().lower()

        if LABEL_SECTION in label:
            in_broiler = True
            continue

        if not in_broiler:
            continue

        if row_feed is None and LABEL_FEED in label:
            row_feed = r
        elif row_market is None and LABEL_MARKET in label:
            row_market = r
        elif row_spread is None and LABEL_PRICE_LESS in label:
            row_spread = r

        if row_feed and row_market and row_spread:
            break

        # Sai do bloco se encontrou outra secao principal
        if row_feed is not None:
            for kw in ("turkey", "cattle", "hog", "pork", "beef", "lamb", "egg (index"):
                if kw in label:
                    in_broiler = False
                    break

    missing = [n for n, r in [("Feed cost", row_feed), ("Market price", row_market),
                                ("Price less cost", row_spread)] if r is None]
    if missing:
        raise RuntimeError(
            f"Nao encontrei no arquivo USDA: {missing}. Layout pode ter mudado."
        )

    print(f"  Linhas USDA: Feed cost={row_feed}, Market price={row_market}, Price less cost={row_spread}")

    result: dict[tuple[int, int], tuple[float, float, float]] = {}
    for col, ym in col_to_ym.items():
        feed   = _cell_to_float(ws.cell(row=row_feed,   column=col).value)
        market = _cell_to_float(ws.cell(row=row_market, column=col).value)
        spread = _cell_to_float(ws.cell(row=row_spread, column=col).value)
        if feed is not None and market is not None and spread is not None:
            result[ym] = (feed, market, spread)

    print(f"  Dados extraidos: {len(result)} meses com valores completos.")
    wb.close()
    return result


# ─────────────────────────────────────────────────────────────────────────────
# 3. Atualizacao do FrangoUS.xlsm via openpyxl
# ─────────────────────────────────────────────────────────────────────────────

def _val_is_empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and v.strip().lower() in ("sem dados", ""):
        return True
    return False


def _approx_equal(a: float, b: float, tol: float = 1e-9) -> bool:
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def update_frangous(usda_data: dict[tuple[int, int], tuple[float, float, float]]) -> dict:
    """
    Abre FrangoUS.xlsm via openpyxl (keep_vba=True) e atualiza Q, R, S.
    """
    # Baixa a planilha do Google Drive (em vez do drive montado em G:\).
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import drive_io
    fid = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(fid)

    wb = openpyxl.load_workbook(str(path), keep_vba=True, data_only=False)
    try:
        ws = wb[SHEET_NAME]

        updates: list[str] = []
        novos   = 0
        revisoes = 0

        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            date_val = ws.cell(row=row, column=COL_DATE).value
            if date_val is None:
                continue
            if isinstance(date_val, (dt.datetime, dt.date)):
                ym = (date_val.year, date_val.month)
            else:
                continue

            usda_vals = usda_data.get(ym)
            if usda_vals is None:
                continue

            feed_new, market_new, spread_new = usda_vals

            for col, new_val, label in [
                (COL_FEED_COST, feed_new,    "Q"),
                (COL_WHOLESALE, market_new,  "R"),
                (COL_SPREAD,    spread_new,  "S"),
            ]:
                cell = ws.cell(row=row, column=col)
                old_val = cell.value

                if _val_is_empty(old_val):
                    cell.value = new_val
                    novos += 1
                    updates.append(
                        f"  [{ym[0]}-{ym[1]:02d}] {label}: '{old_val}' -> {new_val:.4f}"
                    )
                elif isinstance(old_val, (int, float)):
                    if not _approx_equal(float(old_val), new_val):
                        cell.value = new_val
                        revisoes += 1
                        updates.append(
                            f"  [{ym[0]}-{ym[1]:02d}] {label}: {float(old_val):.4f} -> {new_val:.4f} (revisao)"
                        )

        total = novos + revisoes
        if total > 0:
            wb.save(str(path)); drive_io.update_file(fid, path)
            print(f"Salvo. {total} celula(s) alterada(s) ({novos} nova(s), {revisoes} revisao/revisoes).")
            for u in updates[:50]:
                print(u)
            if len(updates) > 50:
                print(f"  ... (+{len(updates) - 50} mais)")
        else:
            print("Nenhuma alteracao necessaria — planilha ja esta atualizada.")

        return {"alteracoes": total, "novos": novos, "revisoes": revisoes}
    finally:
        wb.close()


# ─────────────────────────────────────────────────────────────────────────────
# 4. Main
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] atualiza_usda_frangous.py")
    print(f"Arquivo: {XLSX_NAME}")

    xlsx_bytes = download_usda_xlsx()

    print("Parseando dados USDA...")
    usda_data = parse_usda_broiler_data(xlsx_bytes)

    print(f"Atualizando {XLSX_NAME}...")
    result = update_frangous(usda_data)

    print(
        f"\n[{dt.datetime.now():%Y-%m-%d %H:%M}] Concluido. "
        f"{result['alteracoes']} alteracao(oes) "
        f"({result['novos']} nova(s), {result['revisoes']} revisao/revisoes)."
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit(130)
