"""
atualiza_national_composite_price.py

Atualiza a coluna U ("National Composite", USDc/lb) da aba FrangoUS no FrangoUS.xlsm.

Duas fontes:
  1. USDA ERS "Wholesale Prices" (aba Historical, coluna L = National composite):
     dados mensais oficiais/fechados → valor hardcoded em VERDE (FF00B050).
     O dado estimado do mês corrente (marcado com /* na aba Current) NÃO é usado.

  2. Aba NationalComposite do próprio FrangoUS.xlsm (col D = National Composite Whole Bird):
     usado para calcular a média mensal quando o USDA ainda não publicou o dado
     oficial → fórmula AVERAGEIF em AZUL (FF0070C0).

Lógica de prioridade (roda a cada execução):
  - USDA tem dado fechado E célula vazia/fórmula → hardcode verde (novo ou conversão)
  - USDA tem dado fechado E valor já existe     → compara; atualiza se houve revisão
  - USDA sem dado E NationalComposite tem semanas → fórmula AVERAGEIF azul (se ainda vazia)
  - Fórmula azul já existente → mantém (AVERAGEIF se recalcula ao abrir o Excel)
  - Nenhuma fonte disponível → deixa em branco

Requer: pip install requests openpyxl
"""

from __future__ import annotations

import datetime as dt
import io
import sys
import warnings
from pathlib import Path

import openpyxl
import requests
from openpyxl.styles import Font

sys.stdout.reconfigure(encoding="utf-8")
warnings.filterwarnings("ignore")


# ── Caminhos ──────────────────────────────────────────────────────────────────
DRIVE_ROOT = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
)
XLSX_NAME = "FrangoUS.xlsm"

# ── Planilha ──────────────────────────────────────────────────────────────────
SHEET_FRANGOUS = "FrangoUS"
SHEET_NC       = "NationalComposite"
FIRST_DATA_ROW = 5   # linha 5 em diante (abaixo dos cabeçalhos)
COL_DATE       = 2   # B — datetime mensal
COL_U          = 21  # U — National Composite

# ── Cores ─────────────────────────────────────────────────────────────────────
COLOR_GREEN = "FF00B050"  # dado oficial USDA fechado
COLOR_BLUE  = "FF0070C0"  # fórmula AVERAGEIF (média semanal)

# ── URL ───────────────────────────────────────────────────────────────────────
URL_WHOLESALE = "https://www.ers.usda.gov/media/5538/wholesale-prices.xlsx"


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _averageif_formula(row: int) -> str:
    return f"=AVERAGEIF(NationalComposite!$B:$B,B{row},NationalComposite!$D:$D)"


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _is_empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and v.strip() in ("", "-", "sem dados"):
        return True
    return False


def _approx_equal(a: float, b: float, tol: float = 1e-6) -> bool:
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


# ─────────────────────────────────────────────────────────────────────────────
# 1. Download
# ─────────────────────────────────────────────────────────────────────────────

def localizar_planilha() -> Path:
    """Localiza a planilha sem depender dos nomes das pastas intermediárias."""
    encontrados = list(DRIVE_ROOT.rglob(XLSX_NAME))

    if not encontrados:
        raise FileNotFoundError(
            f"Planilha {XLSX_NAME!r} não encontrada dentro de {DRIVE_ROOT}"
        )

    if len(encontrados) > 1:
        caminhos = "\n".join(f"  - {path}" for path in encontrados)
        raise RuntimeError(
            f"Foram encontradas várias planilhas chamadas {XLSX_NAME!r}:\n"
            f"{caminhos}"
        )

    return encontrados[0]


def download(url: str) -> bytes:
    s = requests.Session()
    s.headers.update({
        "User-Agent": "Mozilla/5.0 (compatible; FrangoUS-updater/1.0)",
        "Accept-Language": "en-US,en;q=0.9",
    })
    print(f"Baixando: {url}")
    resp = s.get(url, timeout=120)
    resp.raise_for_status()
    print(f"  {len(resp.content) / 1024:.1f} KB recebidos.")
    return resp.content


# ─────────────────────────────────────────────────────────────────────────────
# 2. Parse USDA Wholesale Prices — aba Historical
# ─────────────────────────────────────────────────────────────────────────────

def parse_usda_historical(xlsx_bytes: bytes) -> dict[tuple[int, int], float]:
    """
    Lê a aba 'Historical' do arquivo Wholesale Prices do USDA ERS.
    Retorna {(ano, mes): national_composite_cents_lb}.

    Estrutura:
      Linhas 1-4: cabeçalhos
      Col A: Period (datetime), Col L (índice 11): National composite (cents/lb)

    Apenas meses com valor numérico são incluídos; 'NA' é ignorado.
    O mês estimado (Current) não está nessa aba, portanto nunca é capturado.
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb["Historical"]

    result: dict[tuple[int, int], float] = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        date_val = row[0]   # col A
        nc_val   = row[11]  # col L

        if not isinstance(date_val, (dt.datetime, dt.date)):
            continue
        if not isinstance(nc_val, (int, float)):
            continue
        f = float(nc_val)
        if f != f:  # NaN check
            continue

        result[(date_val.year, date_val.month)] = f

    wb.close()
    print(f"  USDA Historical: {len(result)} meses com National Composite.")
    if result:
        last = max(result)
        print(f"  Último dado fechado: {last[0]}-{last[1]:02d}")
    return result


# ─────────────────────────────────────────────────────────────────────────────
# 3. Meses disponíveis no NationalComposite (col B = mês, col D = Whole Bird)
# ─────────────────────────────────────────────────────────────────────────────

def get_nc_months(ws_nc) -> set[tuple[int, int]]:
    """Conjunto de meses que têm ao menos 1 semana na aba NationalComposite."""
    months: set[tuple[int, int]] = set()
    for row in ws_nc.iter_rows(min_row=4, values_only=True):
        mes = row[1]  # col B
        if isinstance(mes, (dt.datetime, dt.date)):
            months.add((mes.year, mes.month))
    return months


# ─────────────────────────────────────────────────────────────────────────────
# 4. Atualização da coluna U
# ─────────────────────────────────────────────────────────────────────────────

def update_frangous(
    path: Path,
    usda_data: dict[tuple[int, int], float],
    nc_months: set[tuple[int, int]],
) -> dict:
    wb = openpyxl.load_workbook(str(path), keep_vba=True, data_only=False)
    try:
        ws = wb[SHEET_FRANGOUS]

        # Fonte de referência: copia nome/tamanho da última célula verde conhecida
        ref = ws.cell(row=319, column=COL_U).font

        def make_font(color: str) -> Font:
            return Font(name=ref.name, size=ref.size,
                        bold=ref.bold, italic=ref.italic, color=color)

        novos_verde  = 0
        novos_azul   = 0
        conv_verde   = 0  # fórmula azul convertida para verde
        revisoes     = 0
        log: list[str] = []

        for row_idx in range(FIRST_DATA_ROW, ws.max_row + 1):
            date_val = ws.cell(row=row_idx, column=COL_DATE).value
            if not isinstance(date_val, (dt.datetime, dt.date)):
                continue

            ym   = (date_val.year, date_val.month)
            cell = ws.cell(row=row_idx, column=COL_U)
            cur  = cell.value
            usda = usda_data.get(ym)

            if usda is not None:
                # — USDA tem dado fechado —
                if _is_formula(cur):
                    # Fórmula azul → substitui por hardcode verde
                    cell.value         = usda
                    cell.font          = make_font(COLOR_GREEN)
                    cell.number_format = "0.00"
                    conv_verde += 1
                    log.append(f"  [{ym[0]}-{ym[1]:02d}] fórmula → {usda:.2f} (verde)")

                elif _is_empty(cur):
                    cell.value         = usda
                    cell.font          = make_font(COLOR_GREEN)
                    cell.number_format = "0.00"
                    novos_verde += 1
                    log.append(f"  [{ym[0]}-{ym[1]:02d}] novo verde: {usda:.2f}")

                elif isinstance(cur, (int, float)):
                    if not _approx_equal(float(cur), usda):
                        old = float(cur)
                        cell.value = usda
                        # mantém o font existente (já deve ser verde)
                        revisoes += 1
                        log.append(f"  [{ym[0]}-{ym[1]:02d}] revisão: {old:.2f} → {usda:.2f}")

            elif ym in nc_months:
                # — Sem dado USDA, mas NationalComposite tem semanas —
                if _is_empty(cur):
                    cell.value         = _averageif_formula(row_idx)
                    cell.font          = make_font(COLOR_BLUE)
                    cell.number_format = "0.00"
                    novos_azul += 1
                    log.append(f"  [{ym[0]}-{ym[1]:02d}] nova fórmula azul (média semanal)")
                # Se já tem fórmula azul, mantém — AVERAGEIF se recalcula no Excel

        total = novos_verde + novos_azul + conv_verde + revisoes
        if total > 0:
            wb.save(str(path))
            print(
                f"Salvo. {total} alteração(ões): "
                f"{novos_verde} verde(s) nova(s), "
                f"{novos_azul} azul(ais) nova(s), "
                f"{conv_verde} fórmula→verde, "
                f"{revisoes} revisão(ões)."
            )
            for entry in log[:50]:
                print(entry)
            if len(log) > 50:
                print(f"  ... (+{len(log) - 50} mais)")
        else:
            print("Nenhuma alteração necessária — planilha já está atualizada.")

        return {
            "total":       total,
            "novos_verde": novos_verde,
            "novos_azul":  novos_azul,
            "conv_verde":  conv_verde,
            "revisoes":    revisoes,
        }
    finally:
        wb.close()


# ─────────────────────────────────────────────────────────────────────────────
# 5. Main
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    # Baixa a planilha do Google Drive (em vez do drive montado em G:\).
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import drive_io
    fid = drive_io.find_file_id(XLSX_NAME)
    path = Path(drive_io.download_to_temp(fid))
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] atualiza_national_composite_price.py")
    print(f"Arquivo: {path.name}\n")

    xlsx_bytes = download(URL_WHOLESALE)
    print()

    print("Parseando dados USDA...")
    usda_data = parse_usda_historical(xlsx_bytes)
    print()

    print("Lendo meses disponíveis na aba NationalComposite...")
    wb_nc = openpyxl.load_workbook(str(path), keep_vba=True, data_only=True)
    ws_nc = wb_nc[SHEET_NC]
    nc_months = get_nc_months(ws_nc)
    wb_nc.close()
    print(f"  {len(nc_months)} mês(es) com dados semanais.")
    if nc_months:
        print(f"  Último mês: {max(nc_months)[0]}-{max(nc_months)[1]:02d}")
    print()

    print(f"Atualizando coluna U (National Composite)...")
    result = update_frangous(path, usda_data, nc_months)
    if result["total"] > 0:
        drive_io.update_file(fid, str(path))   # sobe de volta pro Drive (in-place)

    print(
        f"\n[{dt.datetime.now():%Y-%m-%d %H:%M}] Concluído. "
        f"{result['total']} alteração(ões) "
        f"({result['novos_verde']} verde(s), {result['novos_azul']} azul(ais), "
        f"{result['conv_verde']} fórmula→verde, {result['revisoes']} revisão(ões))."
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit(130)
