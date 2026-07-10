"""
Atualizacao mensal ERS USDA - aba 'Abates' de BeefUS.xlsm.

Versao cloud-capable: baixa o BeefUS.xlsm via Google Drive, escreve valores com
openpyxl e sobe o arquivo de volta. Formulas novas ficam sem cache ate o
Excel/Bloomberg recalcular.
"""

import io
import time
import datetime as dt
import re
import sys
from copy import copy
from pathlib import Path

import requests
import openpyxl
from openpyxl.utils import column_index_from_string

XLSX_NAME = "BeefUS.xlsm"
SHEET_NAME = "Abates"
ANO_INICIO = 2000

USDA_URL = "https://www.ers.usda.gov/media/5529/meat-statistics-tables-historical.xlsx"
TIMEOUT = 120
MAX_RETRIES = 3
RETRY_SLEEP = 5
REVISION_MONTHS = 12

SC_STEERS = 2
SC_HEIFERS = 3
SC_BEEFCOW = 5
SC_DAIRYCOW = 6
SC_BULLS = 7
SC_CALVES_C = 8
SW_CATTLE = 1
SW_CALVES = 2

COL_STEERS = "C"
COL_HEIFERS = "D"
COL_BEEFCOW = "E"
COL_DAIRYCOW = "F"
COL_BULLS = "G"
COL_CALVES = "I"
COL_WT_CATTLE = "P"
COL_WT_CALVES = "Q"
FORMULA_COLS = ["H", "J", "K", "L", "M", "N"]

MONTH_ABBR = {
    "Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
    "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12,
}

_REF = re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?![A-Za-z_(])")


def shift_formula(formula, delta: int):
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(m):
        cabs, col, rabs, row = m.groups()
        return m.group(0) if rabs == "$" else f"{cabs}{col}{rabs}{int(row) + delta}"

    return _REF.sub(repl, formula)


def copy_style(src, dst) -> None:
    dst.number_format = src.number_format
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def copy_row_format_and_formulas(ws, template_row: int, target_row: int) -> None:
    for col in range(1, column_index_from_string("Q") + 1):
        copy_style(ws.cell(template_row, col), ws.cell(target_row, col))
    for col_letter in FORMULA_COLS:
        col = column_index_from_string(col_letter)
        src = ws.cell(template_row, col)
        dst = ws.cell(target_row, col)
        dst.value = shift_formula(src.value, target_row - template_row)
        copy_style(src, dst)


def row_of(ano: int, mes: int) -> int:
    return 3 + (ano - ANO_INICIO) * 12 + (mes - 1)


def parse_usda_date(s):
    if s is None:
        return None
    parts = str(s).strip().split("-")
    if len(parts) != 2 or parts[0].strip() not in MONTH_ABBR:
        return None
    try:
        return int(parts[1].strip()), MONTH_ABBR[parts[0].strip()]
    except ValueError:
        return None


def _diff(a, b) -> bool:
    try:
        return abs(float(a) - float(b)) > 0.01
    except (TypeError, ValueError):
        return a != b


def col_idx(letter: str) -> int:
    return column_index_from_string(letter)


def download_usda() -> bytes:
    last_err = None
    for tentativa in range(1, MAX_RETRIES + 1):
        try:
            print(f"  Baixando USDA Excel (tentativa {tentativa}/{MAX_RETRIES})...")
            r = requests.get(USDA_URL, timeout=TIMEOUT)
            r.raise_for_status()
            print(f"  Download OK - {len(r.content) / 1024:.0f} KB")
            return r.content
        except Exception as e:
            last_err = e
            if tentativa < MAX_RETRIES:
                time.sleep(RETRY_SLEEP)
    raise RuntimeError(f"Download USDA falhou em {MAX_RETRIES} tentativas: {last_err}")


def parse_usda(conteudo: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    counts = {}
    ws_c = wb["SlaughterCounts-Full"]
    for i, row in enumerate(ws_c.iter_rows(values_only=True)):
        if i < 3:
            continue
        parsed = parse_usda_date(row[0])
        if parsed is None:
            continue
        ano, mes = parsed
        if ano < ANO_INICIO:
            continue
        counts[(ano, mes)] = {
            "steers": row[SC_STEERS],
            "heifers": row[SC_HEIFERS],
            "beef_cows": row[SC_BEEFCOW],
            "dairy_cows": row[SC_DAIRYCOW],
            "bulls": row[SC_BULLS],
            "calves": row[SC_CALVES_C],
            "wt_cattle": None,
            "wt_calves": None,
        }

    ws_w = wb["SlaughterWeights-Full"]
    for i, row in enumerate(ws_w.iter_rows(values_only=True)):
        if i < 3:
            continue
        parsed = parse_usda_date(row[0])
        if parsed is None:
            continue
        ano, mes = parsed
        if (ano, mes) in counts:
            counts[(ano, mes)]["wt_cattle"] = row[SW_CATTLE]
            counts[(ano, mes)]["wt_calves"] = row[SW_CALVES]

    print(f"  USDA: {len(counts)} meses mensais lidos ({ANO_INICIO}+)")
    return counts


def last_template_row(ws) -> int | None:
    c = col_idx(COL_STEERS)
    for row in range(ws.max_row, 2, -1):
        if ws.cell(row, c).value is not None:
            return row
    return None


def processar(ws, dados: dict) -> dict:
    hoje = dt.date.today()
    corte_rev = hoje - dt.timedelta(days=REVISION_MONTHS * 31)
    corte_ano_mes = (corte_rev.year, corte_rev.month)
    template_row = last_template_row(ws)

    contadores = {"novas": 0, "revisoes": 0}
    revisoes = []
    linhas_novas = set()
    mudou = False

    for (ano, mes), vals in sorted(dados.items()):
        linha = row_of(ano, mes)
        eh_rev = (ano, mes) >= corte_ano_mes
        rotulo = f"{mes:02d}/{ano}"

        data_b = ws.cell(linha, 2).value
        if isinstance(data_b, dt.datetime):
            data_b = data_b.date()
        if isinstance(data_b, dt.date) and (data_b.year != ano or data_b.month != mes):
            raise RuntimeError(f"Divergencia col B linha {linha}: esperado {rotulo}, encontrado {data_b:%m/%Y}")

        writes = [
            (COL_STEERS, "Steers", vals["steers"]),
            (COL_HEIFERS, "Heifers", vals["heifers"]),
            (COL_BEEFCOW, "Beef Cows", vals["beef_cows"]),
            (COL_DAIRYCOW, "Dairy Cows", vals["dairy_cows"]),
            (COL_BULLS, "Bulls", vals["bulls"]),
            (COL_CALVES, "Calves", vals["calves"]),
            (COL_WT_CATTLE, "Wt Cattle", vals["wt_cattle"]),
            (COL_WT_CALVES, "Wt Calves", vals["wt_calves"]),
        ]

        row_was_empty = all(ws.cell(linha, col_idx(col)).value is None for col, _, _ in writes)
        row_changed = False

        for col_letter, label, novo in writes:
            if novo is None:
                continue
            cell = ws.cell(linha, col_idx(col_letter))
            atual = cell.value
            if atual is None:
                cell.value = novo
                contadores["novas"] += 1
                row_changed = True
            elif eh_rev and _diff(atual, novo):
                revisoes.append((rotulo, label, atual, novo))
                contadores["revisoes"] += 1
                cell.value = novo
                row_changed = True

        if row_changed:
            mudou = True
            if row_was_empty:
                linhas_novas.add(linha)
                if template_row is not None and template_row != linha:
                    copy_row_format_and_formulas(ws, template_row, linha)
                    for col_letter, _, novo in writes:
                        if novo is not None:
                            ws.cell(linha, col_idx(col_letter)).value = novo

    return {
        "salvar": mudou,
        "contadores": contadores,
        "revisoes": revisoes,
        "linhas_novas": sorted(linhas_novas),
        "template_row": template_row,
        "meses_dados": len(dados),
    }


def print_relatorio(r: dict):
    print("\n== Atualizacao USDA - BeefUS Abates ==")
    print(f"Meses processados : {r['meses_dados']} ({ANO_INICIO}+)")
    print(f"Linhas novas      : {len(r['linhas_novas'])}")
    if r["linhas_novas"]:
        print(f"  (template usada : linha {r['template_row']})")
    print(f"Celulas novas     : {r['contadores']['novas']}")
    print(f"Revisoes          : {r['contadores']['revisoes']}")
    for mes, label, antigo, novo in r["revisoes"]:
        print(f"  - {mes} {label}: {antigo} -> {novo}")
    print(f"\nStatus: {'SUCESSO' if r['salvar'] else 'SEM ALTERACOES'}")


def main():
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import drive_io

    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] Iniciando USDA - BeefUS Abates...")
    print("\n[1/3] Baixando e processando Excel do USDA...")
    dados = parse_usda(download_usda())
    if not dados:
        print("Nenhum dado encontrado. Encerrando.")
        return

    print("\n[2/3] Baixando BeefUS.xlsm do Drive...")
    fid = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(fid)

    print("\n[3/3] Gravando dados na planilha...")
    wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
    try:
        resultado = processar(wb[SHEET_NAME], dados)
        if resultado["salvar"]:
            wb.save(path)
            drive_io.update_file(fid, path)
            print("Arquivo salvo no Drive.")
        else:
            print("Nada alterado - arquivo nao salvo.")
    finally:
        wb.close()

    print_relatorio(resultado)


if __name__ == "__main__":
    main()
