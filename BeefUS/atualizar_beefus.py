"""
atualizar_beefus.py (versao openpyxl / nuvem)

Consolidador do BeefUS. Estende formulas e formatos sem Excel, gravando o
BeefUS.xlsm via Google Drive (drive_io). As formulas novas so ganham cache de
valor depois do Excel/Bloomberg recalcular a planilha.
"""

from __future__ import annotations

import datetime as dt
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl

XLSX_NAME = "BeefUS.xlsm"
SHEET_MAIN = "BeefUS"
SHEET_ABATES = "Abates"

MAIN_FIRST_DATA_ROW = 5
ABATES_FIRST_DATA_ROW = 3
ROW_OFFSET_ABATES_TO_MAIN = MAIN_FIRST_DATA_ROW - ABATES_FIRST_DATA_ROW

ABATES_MAIN_COL_START = 4
ABATES_MAIN_COL_END = 10
COL_DATE_MAIN = 2
COL_DATE_ABATES = 2
ABATES_SOURCE_COL_START = 3
ABATES_SOURCE_COL_END = 14
COL_M = 13
COL_N = 14
COL_O = 15
COL_P = 16
COL_Q = 17

_REF = re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?![A-Za-z_(])")


def shift_formula(formula, delta: int):
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(m):
        col_abs, col, row_abs, row = m.groups()
        if row_abs == "$":
            return m.group(0)
        return f"{col_abs}{col}{row_abs}{int(row) + delta}"

    return _REF.sub(repl, formula)


def is_usable_price_formula(value) -> bool:
    """Aceita formulas reais e ignora o placeholder =\"Sem Dados\"."""
    if not isinstance(value, str) or not value.lstrip().startswith("="):
        return False
    body = value.lstrip()[1:].strip()
    return body.casefold() not in {'"sem dados"', "'sem dados'"}


def show_sem_dados_on_error(formula):
    """Mantem a formula original, mas troca erros de calculo pelo texto de fallback."""
    if not is_usable_price_formula(formula):
        return formula
    if formula.lstrip()[1:].upper().startswith("IFERROR("):
        return formula
    return f'=IFERROR({formula.lstrip()[1:]},"Sem Dados")'


def copy_cell(src, dst, *, shift_delta: int = 0, protect_price_formula: bool = False) -> None:
    src_value = src.value
    value = shift_formula(src_value, shift_delta) if isinstance(src_value, str) and src_value.startswith("=") else src_value
    dst.value = show_sem_dados_on_error(value) if protect_price_formula else value
    dst.number_format = src.number_format
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def extend_section(
    ws,
    src_row: int,
    col_ranges: list[tuple[int, int]],
    dst_end: int,
    *,
    protect_price_formula: bool = False,
) -> int:
    if dst_end <= src_row:
        return 0
    for c0, c1 in col_ranges:
        for col in range(c0, c1 + 1):
            src = ws.cell(row=src_row, column=col)
            for dst_row in range(src_row + 1, dst_end + 1):
                copy_cell(
                    src,
                    ws.cell(row=dst_row, column=col),
                    shift_delta=dst_row - src_row,
                    protect_price_formula=protect_price_formula,
                )
    return dst_end - src_row


def last_value_row(ws, col: int, start: int) -> int:
    for row in range(ws.max_row, start - 1, -1):
        value = ws.cell(row=row, column=col).value
        if value not in (None, ""):
            return row
    return start - 1


def last_formula_row(ws, col: int, start: int) -> int:
    for row in range(ws.max_row, start - 1, -1):
        value = ws.cell(row=row, column=col).value
        if isinstance(value, str) and value.startswith("="):
            return row
    return start - 1


def last_usable_price_formula_row(ws, col: int, start: int) -> int:
    for row in range(ws.max_row, start - 1, -1):
        if is_usable_price_formula(ws.cell(row=row, column=col).value):
            return row
    return start - 1


def year_month(value) -> tuple[int, int] | None:
    if isinstance(value, dt.datetime):
        value = value.date()
    if isinstance(value, dt.date):
        return value.year, value.month

    if not isinstance(value, str):
        return None
    match = re.fullmatch(r"\s*(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)[-/](\d{2,4})\s*", value.casefold())
    if not match:
        return None
    month_names = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
    year = int(match.group(2))
    return (2000 + year if year < 100 else year), month_names.index(match.group(1)) + 1


def rows_by_year_month(ws, date_col: int, start: int) -> dict[tuple[int, int], list[int]]:
    rows: dict[tuple[int, int], list[int]] = {}
    for row in range(start, ws.max_row + 1):
        period = year_month(ws.cell(row=row, column=date_col).value)
        if period:
            rows.setdefault(period, []).append(row)
    return rows


def nearest_formula_row(ws, col: int, target_row: int, start: int) -> int | None:
    for row in range(target_row - 1, start - 1, -1):
        value = ws.cell(row=row, column=col).value
        if isinstance(value, str) and value.startswith("="):
            return row
    for row in range(target_row + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=col).value
        if isinstance(value, str) and value.startswith("="):
            return row
    return None


def row_for_year_month(ws, year: int, month: int) -> int | None:
    for row in range(MAIN_FIRST_DATA_ROW, ws.max_row + 1):
        value = ws.cell(row=row, column=COL_DATE_MAIN).value
        if isinstance(value, dt.datetime):
            value = value.date()
        if isinstance(value, dt.date) and value.year == year and value.month == month:
            return row
    return None


def sincronizar_abates(ws_main, ws_abates) -> dict:
    """Preenche somente D:J vazias, associando cada linha pelo mes correspondente."""
    main_rows = rows_by_year_month(ws_main, COL_DATE_MAIN, MAIN_FIRST_DATA_ROW)
    completed_rows: set[int] = set()
    cells_filled = 0
    unmatched_periods = 0

    for abates_row in range(ABATES_FIRST_DATA_ROW, ws_abates.max_row + 1):
        has_source_data = any(
            ws_abates.cell(row=abates_row, column=col).value not in (None, "")
            for col in range(ABATES_SOURCE_COL_START, ABATES_SOURCE_COL_END + 1)
        )
        if not has_source_data:
            continue

        period = year_month(ws_abates.cell(row=abates_row, column=COL_DATE_ABATES).value)
        if not period:
            continue
        candidates = main_rows.get(period, [])
        target_row = next(
            (
                row
                for row in candidates
                if any(
                    ws_main.cell(row=row, column=col).value in (None, "")
                    for col in range(ABATES_MAIN_COL_START, ABATES_MAIN_COL_END + 1)
                )
            ),
            None,
        )
        if target_row is None:
            if not candidates:
                unmatched_periods += 1
            continue

        for col in range(ABATES_MAIN_COL_START, ABATES_MAIN_COL_END + 1):
            target = ws_main.cell(row=target_row, column=col)
            if target.value not in (None, ""):
                continue
            template_row = nearest_formula_row(ws_main, col, target_row, MAIN_FIRST_DATA_ROW)
            if template_row is None:
                continue
            copy_cell(
                ws_main.cell(row=template_row, column=col),
                target,
                shift_delta=target_row - template_row,
            )
            cells_filled += 1
            completed_rows.add(target_row)

    if not cells_filled:
        return {"mudou": False, "msg": "Abates: nenhuma celula vazia para preencher."}

    warning = f"; {unmatched_periods} mes(es) sem data correspondente" if unmatched_periods else ""
    return {
        "mudou": True,
        "msg": f"Abates: {cells_filled} formula(s) reposta(s) em {len(completed_rows)} mes(es){warning}.",
    }

def ajustar_precos(ws_main, today: dt.date) -> dict:
    row_cur = row_for_year_month(ws_main, today.year, today.month)
    if row_cur is None:
        return {"mudou": False, "msg": "Precos: nao encontrei a linha do mes atual na col B."}

    prev = today.replace(day=1) - dt.timedelta(days=1)
    row_prev = row_for_year_month(ws_main, prev.year, prev.month) or row_cur - 1
    if row_prev < MAIN_FIRST_DATA_ROW:
        return {"mudou": False, "msg": "Precos: sem linha anterior para usar como template."}

    mudou = False
    acoes: list[str] = []

    for col, label in [(COL_M, "M"), (COL_Q, "Q")]:
        last_row = last_usable_price_formula_row(ws_main, col=col, start=MAIN_FIRST_DATA_ROW)
        if last_row >= MAIN_FIRST_DATA_ROW and last_row < row_cur:
            extend_section(
                ws_main,
                last_row,
                [(col, col)],
                row_cur,
                protect_price_formula=True,
            )
            mudou = True
            acoes.append(f"{label}: formula copiada linhas {last_row + 1}->{row_cur} ({row_cur - last_row} mes(es))")

        cell = ws_main.cell(row=row_cur, column=col)
        if is_usable_price_formula(cell.value):
            protected = show_sem_dados_on_error(cell.value)
            if protected != cell.value:
                cell.value = protected
                mudou = True
                acoes.append(f"{label}: formula protegida contra erros na linha {row_cur}")

    n_val = ws_main.cell(row=row_cur, column=COL_N).value
    if n_val not in (None, ""):
        for col, label in [(COL_O, "O"), (COL_P, "P")]:
            cell = ws_main.cell(row=row_cur, column=col)
            has_formula = isinstance(cell.value, str) and cell.value.startswith("=")
            if not has_formula and cell.value in (None, ""):
                copy_cell(ws_main.cell(row=row_prev, column=col), cell, shift_delta=row_cur - row_prev)
                mudou = True
                acoes.append(f"{label}: formula copiada {row_prev}->{row_cur} (N preenchida)")

    if not mudou:
        return {"mudou": False, "msg": "Precos: nada para ajustar."}
    return {"mudou": True, "msg": "Precos: " + "; ".join(acoes)}


def main() -> None:
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import drive_io

    fid = drive_io.find_file_id(XLSX_NAME)
    path = drive_io.download_to_temp(fid)
    today = dt.date.today()
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] BeefUS consolidacao - {today:%d/%m/%Y}")

    wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
    try:
        ws_main = wb[SHEET_MAIN]
        ws_abates = wb[SHEET_ABATES]

        r1 = sincronizar_abates(ws_main, ws_abates)
        print(" - " + r1["msg"])
        r2 = ajustar_precos(ws_main, today)
        print(" - " + r2["msg"])

        if r1.get("mudou") or r2.get("mudou"):
            wb.save(path)
            drive_io.update_file(fid, path)
            print("Salvo. Formulas novas serao calculadas no PC Bloomberg.")
        else:
            print("Nada alterado - nao salvou.")
    finally:
        wb.close()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit(130)
