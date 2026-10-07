"""Leitura das tabelas de abate do USDA/ERS (meat-statistics-tables-historical).

POR QUE EXISTE
--------------
Em outubro de 2026 o USDA mexeu na planilha e quebrou BeefUS e FrangoUS de
duas formas ao mesmo tempo:

1. As abas perderam o sufixo: "SlaughterCounts-Full" virou "SlaughterCounts",
   "SlaughterWeights-Full" virou "SlaughterWeights". Isso dava KeyError, que
   pelo menos falha alto.

2. Cada bloco da SlaughterCounts perdeu uma coluna. As posicoes fixas que os
   scripts usavam passaram a apontar para o vizinho:

       --Beef cows   era 5   virou 4
       --Dairy cows  era 6   virou 5
       Broilers      era 31  virou 29   (o 31 agora e Turkeys)

   Esse segundo caso e o perigoso: so renomear a aba faria o FrangoUS gravar
   PERU como se fosse frango, sem erro nenhum.

Por isso aqui as colunas sao resolvidas pelo TEXTO do cabecalho. Se o USDA
mexer de novo na ordem, nada quebra; se apagar ou renomear uma coluna, quebra
alto, com o nome do que faltou.
"""
from __future__ import annotations

import io

import openpyxl

LINHA_CABECALHO = 2  # 0-based: linha 3 da planilha
PRIMEIRA_LINHA_DADOS = 3


def abrir(conteudo: bytes):
    return openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)


def aba(wb, *nomes: str):
    """Primeira aba que existir entre os nomes dados."""
    for nome in nomes:
        if nome in wb.sheetnames:
            return wb[nome]
    raise RuntimeError(
        f"Nenhuma destas abas existe na planilha do USDA: {', '.join(nomes)}. "
        f"Abas disponiveis: {', '.join(wb.sheetnames)}."
    )


def cabecalho(ws) -> list[str]:
    for i, linha in enumerate(ws.iter_rows(values_only=True)):
        if i == LINHA_CABECALHO:
            return [("" if c is None else str(c).strip()) for c in linha]
    return []


def colunas(ws, *nomes: str) -> dict[str, int]:
    """{nome do cabecalho: indice}. Usa a PRIMEIRA ocorrencia de cada nome.

    A SlaughterCounts repete os mesmos rotulos em dois blocos (Commercial e
    Federally inspected). Os scripts sempre leram o primeiro, o Commercial --
    menos Broilers, que so existe no segundo. Pegar a primeira ocorrencia
    preserva os dois comportamentos.
    """
    linha = cabecalho(ws)
    achados: dict[str, int] = {}
    for nome in nomes:
        try:
            achados[nome] = linha.index(nome)
        except ValueError:
            raise RuntimeError(
                f"Coluna {nome!r} nao encontrada na aba {ws.title!r} do USDA. "
                f"Cabecalho atual: {[c for c in linha if c]}"
            )
    return achados
