"""Atualizacao SIDRA 1092 - aba Abates de BeefBR.xlsm, via Drive."""
from __future__ import annotations
import datetime as dt
import sys
from copy import copy
from pathlib import Path
import openpyxl
import requests

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import drive_io

XLSX_NAME="BeefBR.xlsm"; SHEET_NAME="Abates"
URL="https://apisidra.ibge.gov.br/values/t/1092/n1/1/v/284,285/p/{periods}/c12716/115233,115234,115235/c18/55,56,111734,111735,57/c12529/118225"
COLS={"55":(3,15),"56":(4,16),"111734":(5,17),"111735":(6,18),"57":(7,19)}
MONTH={"115233":1,"115234":2,"115235":3}

def row_of(y,m): return 3+(y-1997)*12+m-1

def shift(formula,delta):
 import re
 if not isinstance(formula,str) or not formula.startswith('='): return formula
 return re.sub(r'(?<![A-Za-z0-9_])(\$?[A-Za-z]{1,3})(\$?)(\d+)',lambda x:x.group(0) if x.group(2)=='$' else f'{x.group(1)}{x.group(2)}{int(x.group(3))+delta}',formula)

def recent_periods():
 y,q=dt.date.today().year,(dt.date.today().month-1)//3+1
 result=[]
 for _ in range(8):
  key=f'{y}{q:02d}'
  try:
   if requests.get(URL.format(periods=key),timeout=30).json()[1:]: break
  except Exception: pass
  q-=1
  if q==0:y,q=y-1,4
 for _ in range(5):
  result.append(f'{y}{q:02d}');q-=1
  if q==0:y,q=y-1,4
 return list(reversed(result))

def fetch():
 r=requests.get(URL.format(periods=','.join(recent_periods())),timeout=90);r.raise_for_status();out={}
 for item in r.json()[1:]:
  var,cat,mc,p=item.get('D2C'),item.get('D5C'),item.get('D4C'),item.get('D3C')
  if var not in {'284','285'} or cat not in COLS or mc not in MONTH or not p:continue
  raw=str(item.get('V','')).strip()
  if raw=='X':continue
  try:value=0 if raw in {'','...','-'} else float(raw)
  except ValueError:continue
  y,q=int(p[:4]),int(p[4:]);m=(q-1)*3+MONTH[mc]
  out.setdefault((y,m,cat),{})['heads' if var=='284' else 'weight']=value
 return out

def copy_formula(ws,src,dst,cols):
 for col in cols:
  s,t=ws.cell(src,col),ws.cell(dst,col);t._style=copy(s._style);t.number_format=s.number_format;t.value=shift(s.value,dst-src)

def main():
 data=fetch();fid=drive_io.find_file_id(XLSX_NAME);path=drive_io.download_to_temp(fid)
 wb=openpyxl.load_workbook(path,keep_vba=True,data_only=False)
 try:
  ws=wb[SHEET_NAME];template=next((r for r in range(ws.max_row,2,-1) if any(ws.cell(r,c).value not in (None,'') for c in range(3,8))),None);changes=0
  for (y,m,cat),values in data.items():
   row=row_of(y,m)
   if row>ws.max_row:continue
   new=all(ws.cell(row,c).value in (None,'') for c in range(3,8))
   if new and template:copy_formula(ws,template,row,tuple(range(8,14))+(20,))
   a,b=COLS[cat]
   for key,col in (('heads',a),('weight',b)):
    if key in values and ws.cell(row,col).value!=values[key]:ws.cell(row,col).value=values[key];changes+=1
  if changes:wb.save(path);drive_io.update_file(fid,path)
  print(f'SIDRA BeefBR: {changes} celula(s) atualizada(s).')
 finally:wb.close()
if __name__=='__main__':main()
