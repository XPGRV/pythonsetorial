"""Atualiza SECEX (EUA) de BeefBR.xlsm pelo USDA GATS, via Drive."""
from __future__ import annotations
import datetime as dt,sys
from copy import copy
from pathlib import Path
import openpyxl,requests
from bs4 import BeautifulSoup
sys.path.insert(0,str(Path(__file__).resolve().parent.parent));import drive_io
XLSX_NAME='BeefBR.xlsm';SHEET='SECEX';BASE='https://apps.fas.usda.gov/GATS/';PRODUCTS=(('suina','FATUS','1200AG'),('bovina','FATUS','1130AG'),('frango','FAS','0504AT'));COLS={18:('suina','q'),19:('bovina','q'),20:('frango','q'),22:('suina','p'),23:('bovina','p'),24:('frango','p')}
def month(v):
 if isinstance(v,dt.datetime):v=v.date()
 return (v.year,v.month) if isinstance(v,dt.date) else None
def fields(soup):
 out={};form=soup.find('form') or soup
 for e in form.find_all(['input','select']):
  n=e.get('name')
  if not n:continue
  if e.name=='select':o=e.find('option',selected=True) or e.find('option');out[n]=o.get('value','') if o else ''
  elif e.get('type','text').lower() in {'hidden','text'}:out[n]=e.get('value','')
 return out
def table_data(t):
 out={}
 for tr in t.find_all('tr'):
  c=tr.find_all('td')
  if len(c)!=39 or not c[9].get_text(strip=True)[:4].isdigit():continue
  y=int(c[9].get_text(strip=True)[:4])
  for m in range(1,13):
   try:out[(y,m)]={'p':float(c[11+(m-1)*2].get_text(strip=True).replace(',','')),'q':float(c[12+(m-1)*2].get_text(strip=True).replace(',',''))}
   except ValueError:pass
 return out
def fetch():
 s=requests.Session();s.headers['User-Agent']='Mozilla/5.0';s.get(BASE+'detectscreen.aspx?action=set&res=1920x1080&d=24&pw=1920',timeout=30)
 home=BeautifulSoup(s.get(BASE+'default.aspx',timeout=30).text,'html.parser');a=home.find('a',string=lambda x:x and 'Standard Query' in x);url=BASE+(a['href'].lstrip('/') if a else 'ExpressQuery1.aspx');s0=BeautifulSoup(s.get(url,timeout=30).text,'html.parser');action=(s0.find('form') or {}).get('action','');post=action if action.startswith('http') else BASE+action.lstrip('/');pre='ctl00$ContentPlaceHolder1$';out={}
 for n,g,p in PRODUCTS:
  one=fields(s0);one.update({'__EVENTTARGET':pre+'ddlProductGroup',pre+'ddlProductGroup':g});s1=BeautifulSoup(s.post(post,data=one,timeout=60).text,'html.parser');q=fields(s1);q.update({pre+'ddlProductGroup':g,pre+'lb_Products':p,pre+'lb_Partners':'R00',pre+'ddlDateSeries':'Monthly',pre+'ddlStartYear':'2000',pre+'ddlEndYear':str(dt.date.today().year),pre+'ddlValueType':'UVAL',pre+'ddlValueUnit':'D',pre+'btnRetrieveData':'Retrieve Data'});t=BeautifulSoup(s.post(post,data=q,timeout=120).text,'html.parser').find('table',id=lambda x:x and 'GridView' in x);out[n]=table_data(t) if t else {}
 return out
def main():
 data=fetch();fid=drive_io.find_file_id(XLSX_NAME);path=drive_io.download_to_temp(fid);wb=openpyxl.load_workbook(path,keep_vba=True,data_only=False)
 try:
  ws=wb[SHEET];changes=0
  for r in range(3,ws.max_row+1):
   p=month(ws.cell(r,2).value)
   if not p:continue
   for col,(prod,key) in COLS.items():
    cell=ws.cell(r,col);val=data.get(prod,{}).get(p,{}).get(key)
    if cell.data_type=='f' or cell.value not in (None,'') or val is None:continue
    cell.value=round(val/1000,4 if key=='q' else 6)
    if r>3:cell._style=copy(ws.cell(r-1,col)._style)
    changes+=1
  if changes:wb.save(path);drive_io.update_file(fid,path)
  print(f'USDA GATS BeefBR: {changes} celula(s) atualizada(s).')
 finally:wb.close()
if __name__=='__main__':main()
