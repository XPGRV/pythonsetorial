#!/usr/bin/env python3
"""
IMEA — Extractor de FRETE DE GRAOS via API (/api/seriehistorica).

Reaproveita o mecanismo do repositorio pedronipalhares/IMEA (imea_extractor.py):
autenticacao OAuth2 (/token) -> chamada ao endpoint generico /api/seriehistorica
com o indicador de frete, quebrando o historico em janelas mensais para driblar
o teto de ~10 registros por requisicao, com concorrencia.

Descoberto a partir do payload real da pagina "Serie historica" do portal:
  - indicador de frete: 708192508843130882  ("Preco disponivel do Frete de Graos")
  - tipolocalidade / tipoDestino = "3" (nivel CIDADE)
  - safra = [] (frete nao e amarrado a safra)
  - exemplo de rota: estado 51 (MT) cidade 7602 -> estadoDestino 35 (SP) cidade 48500

CREDENCIAIS (nunca hardcode o token):
  export IMEA_USERNAME="seu_login"
  export IMEA_PASSWORD="sua_senha"
  # ou coloque num arquivo .env e use python-dotenv

Dependencias:
  pip install requests pandas python-dateutil python-dotenv openpyxl
"""

import os
import ssl
import copy
import argparse
import logging
import requests
import urllib3
import pandas as pd
import openpyxl
from pathlib import Path
from datetime import date, datetime
from dateutil.relativedelta import relativedelta
from concurrent.futures import ThreadPoolExecutor, as_completed

# Carrega sempre o .env ao lado deste script. O fallback evita depender de
# python-dotenv na maquina que executa o agendamento.
ENV_PATH = Path(__file__).resolve().parent / ".env"
try:
    from dotenv import load_dotenv
    load_dotenv(dotenv_path=ENV_PATH)
except ImportError:
    if ENV_PATH.exists():
        for raw_line in ENV_PATH.read_text(encoding="utf-8-sig").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key:
                os.environ.setdefault(key, value)

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
log = logging.getLogger("imea_frete")

# ---------------------------------------------------------------------------
# CONFIGURACAO
# ---------------------------------------------------------------------------
BASE_URL = "https://api1.imea.com.br"
FRETE_INDICADOR = "708192508843130882"      # "Preco disponivel do Frete de Graos"
TIPO_LOCALIDADE = ["3"]                       # 3 = Cidade (origem e destino)

DATA_INICIO = "2023-01-01"                    # primeiras observacoes das rotas: 13/01/2023
DATA_FIM = date.today().isoformat()
JANELA = "mensal"                             # "mensal" ou "semanal" (use semanal se aparecer aviso de teto)
CAP_SERVIDOR = 10                             # teto de registros por requisicao observado na API
MAX_WORKERS = 10

# Rotas a coletar. Cada rota = origem (estado/cidade) -> destino (estadoDestino/cidadeDestino).
# Codigos de estado sao IBGE (51=MT, 35=SP, 41=PR...). Codigos de cidade sao internos do IMEA
# (pegue-os abrindo os dropdowns Cidade / Cidade Destino e olhando a resposta no Network).
#
# Para coletar TODAS as rotas de uma vez, use MODO_TODAS_ROTAS = True (veja abaixo).
ROTAS = [
    {
        "nome": "Sorriso__Santos",           # Sorriso(MT) cidade 7925 -> Santos(SP) cidade 48500
        "estado": ["51"], "cidade": ["7925"],
        "estadoDestino": ["35"], "cidadeDestino": ["48500"],
    },
    {
        "nome": "Sorriso__Rondonopolis",     # Sorriso(MT) 7925 -> Rondonopolis(MT) 7602
        "estado": ["51"], "cidade": ["7925"],
        "estadoDestino": ["51"], "cidadeDestino": ["7602"],
    },
    {
        "nome": "Rondonopolis__Santos",      # Rondonopolis(MT) 7602 -> Santos(SP) 48500
        "estado": ["51"], "cidade": ["7602"],
        "estadoDestino": ["35"], "cidadeDestino": ["48500"],
    },
    # adicione mais rotas aqui, ex.:
    # {"nome": "Sorriso__Paranagua", "estado":["51"],"cidade":["7925"],
    #  "estadoDestino":["41"],"cidadeDestino":["<id>"]},
]

# Se True, ignora ROTAS e consulta o indicador sem filtro de origem/destino (todas as rotas).
# Nesse caso, use JANELA = "semanal" para nao estourar o teto de 10 registros por requisicao.
MODO_TODAS_ROTAS = False

WORKBOOK_PATH = Path(
    r"G:\.shortcut-targets-by-id\11q1ngdqbySnDMCQSP4fiwPztT6ju2F4Z"
    r"\Arquivos\Setorial - Database\Transportes.xlsm"
)
SHEET_NAME = "Preço Fretes"
FIRST_DATA_ROW = 4

COL_QUARTER = 1
COL_MES = 2
COL_SEMANA = 3
ROUTE_COLUMNS = {
    "Sorriso__Santos": 4,
    "Rondonopolis__Santos": 5,
    "Sorriso__Rondonopolis": 6,
}


# ---------------------------------------------------------------------------
# AUTENTICACAO (OAuth2 password grant) — token vem de env, nunca hardcoded
# ---------------------------------------------------------------------------
class TLSAdapter(requests.adapters.HTTPAdapter):
    """Adapter que relaxa a verificacao TLS (o cert da api1.imea.com.br e atipico)."""
    def init_poolmanager(self, *args, **kwargs):
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
        kwargs["ssl_context"] = ctx
        return super().init_poolmanager(*args, **kwargs)


def make_session_and_token():
    user = os.getenv("IMEA_USERNAME")
    pwd = os.getenv("IMEA_PASSWORD")
    if not user or not pwd:
        raise SystemExit("Defina IMEA_USERNAME e IMEA_PASSWORD (env ou .env).")

    session = requests.Session()
    session.mount("https://", TLSAdapter())

    log.info("Autenticando na API do IMEA...")
    resp = session.post(
        f"{BASE_URL}/token",
        headers={"content-type": "application/x-www-form-urlencoded",
                 "authorization": "bearer undefined"},
        data={"username": user, "password": pwd, "grant_type": "password", "client_id": "2"},
        verify=False, timeout=30,
    )
    resp.raise_for_status()
    token = resp.json().get("access_token")
    if not token:
        raise SystemExit("Autenticacao sem access_token na resposta.")
    log.info("Autenticado com sucesso.")
    return session, token


# ---------------------------------------------------------------------------
# JANELAS DE DATA
# ---------------------------------------------------------------------------
def gerar_janelas(inicio_str, fim_str, modo):
    inicio = datetime.fromisoformat(inicio_str).date()
    fim = datetime.fromisoformat(fim_str).date()
    passo = relativedelta(months=1) if modo == "mensal" else relativedelta(weeks=1)
    janelas, cur = [], inicio
    while cur <= fim:
        prox = cur + passo
        janelas.append((cur.isoformat(), min(prox - relativedelta(days=1), fim).isoformat()))
        cur = prox
    return janelas


# ---------------------------------------------------------------------------
# PAYLOAD + REQUISICAO
# ---------------------------------------------------------------------------
def build_payload(rota, inicio, fim):
    return {
        "pageSize": 1000,
        "inicio": inicio, "fim": fim,
        "cadeia": [], "grupo": [],
        "indicador": [FRETE_INDICADOR],
        "tipolocalidade": TIPO_LOCALIDADE,
        "estado": rota.get("estado", []),
        "regiao": rota.get("regiao", []),
        "cidade": rota.get("cidade", []),
        "tipoDestino": TIPO_LOCALIDADE,
        "estadoDestino": rota.get("estadoDestino", []),
        "regiaoDestino": rota.get("regiaoDestino", []),
        "cidadeDestino": rota.get("cidadeDestino", []),
        "safra": [],
    }


def fetch_janela(session, token, rota, inicio, fim):
    headers = {
        "accept": "application/json, text/plain, */*",
        "authorization": f"bearer {token}",
        "content-type": "application/json;charset=UTF-8",
        "origin": "https://portal.imea.com.br",
        "referer": "https://portal.imea.com.br/",
        "user-agent": "Mozilla/5.0 (pipeline interno de analise)",
    }
    try:
        r = session.post(f"{BASE_URL}/api/seriehistorica",
                         headers=headers, json=build_payload(rota, inicio, fim),
                         verify=False, timeout=25)
        if r.status_code != 200:
            return {"ok": False, "rota": rota["nome"], "janela": f"{inicio}..{fim}",
                    "erro": f"HTTP {r.status_code}"}
        data = r.json()
        recs = data if isinstance(data, list) else (data.get("data", []) if isinstance(data, dict) else [])
        for rec in recs:
            rec["_rota"] = rota["nome"]
            rec["_janela"] = f"{inicio}..{fim}"
        return {"ok": True, "rota": rota["nome"], "janela": f"{inicio}..{fim}",
                "recs": recs, "n": len(recs), "truncado": len(recs) >= CAP_SERVIDOR}
    except Exception as e:
        return {"ok": False, "rota": rota["nome"], "janela": f"{inicio}..{fim}", "erro": str(e)}


# ---------------------------------------------------------------------------
# ATUALIZACAO DA PLANILHA
# ---------------------------------------------------------------------------
def quarter_str(d):
    return f"{((d.month - 1) // 3) + 1}Q{str(d.year)[2:]}"


def copy_row_format(ws, src_row, dst_row, n_cols=6):
    for col in range(1, n_cols + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.number_format = src.number_format
        dst.alignment = copy.copy(src.alignment)
        dst.border = copy.copy(src.border)
        dst.protection = copy.copy(src.protection)


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        parsed = pd.to_datetime(value, errors="coerce")
        if not pd.isna(parsed):
            return parsed.date()
    return None


def records_by_week(df):
    result = {}
    for rec in df.to_dict("records"):
        week = as_date(rec.get("Data") or rec.get("data") or rec.get("DataReferencia"))
        route = rec.get("_rota")
        if week is None or route not in ROUTE_COLUMNS:
            continue
        raw_value = rec.get("valor2")
        if raw_value is None or pd.isna(raw_value):
            raw_value = rec.get("Valor")
        try:
            value = float(str(raw_value).replace(",", "."))
        except (TypeError, ValueError):
            continue
        result.setdefault(week, {})[route] = value
    return result


def update_workbook(df, dry_run=False):
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Planilha nao encontrada: {WORKBOOK_PATH}")

    weekly = records_by_week(df)
    if not weekly:
        raise RuntimeError("Nenhum registro valido para gravar na planilha.")

    wb = openpyxl.load_workbook(WORKBOOK_PATH, keep_vba=True, data_only=False)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise RuntimeError(
                f"Aba {SHEET_NAME!r} nao encontrada. Abas: {', '.join(wb.sheetnames)}"
            )
        ws = wb[SHEET_NAME]
        original_max_row = ws.max_row

        rows_by_week = {}
        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            week = as_date(ws.cell(row=row, column=COL_SEMANA).value)
            if week is not None:
                rows_by_week[week] = row

        def next_empty_row():
            for row in range(FIRST_DATA_ROW, ws.max_row + 2):
                if ws.cell(row=row, column=COL_SEMANA).value is None:
                    return row
            return ws.max_row + 1

        new_rows = 0
        updated_values = 0
        log_rows = []

        for week in sorted(weekly):
            route_values = weekly[week]
            row = rows_by_week.get(week)
            is_new = row is None
            if is_new:
                row = next_empty_row()
                if row > original_max_row and row > FIRST_DATA_ROW:
                    copy_row_format(ws, row - 1, row)
                new_rows += 1
                rows_by_week[week] = row

            if not dry_run:
                ws.cell(row=row, column=COL_QUARTER).value = quarter_str(week)
                ws.cell(row=row, column=COL_MES).value = week.replace(day=1)
                ws.cell(row=row, column=COL_SEMANA).value = week
                ws.cell(row=row, column=COL_MES).number_format = "[$-416]mmm-yy;@"
                ws.cell(row=row, column=COL_SEMANA).number_format = "dd-mmm"

            changed_routes = []
            for route, value in route_values.items():
                col = ROUTE_COLUMNS[route]
                old = ws.cell(row=row, column=col).value
                if old is None or abs(float(old) - value) > 1e-9:
                    updated_values += 1
                    changed_routes.append(route)
                    if not dry_run:
                        ws.cell(row=row, column=col).value = value
                        ws.cell(row=row, column=col).number_format = "0.00"

            if is_new or changed_routes:
                log_rows.append(
                    f"  {week.isoformat()} | "
                    + " | ".join(
                        f"{route}={route_values[route]:.2f}"
                        for route in ROUTE_COLUMNS if route in route_values
                    )
                )

        if not dry_run and (new_rows or updated_values):
            wb.save(WORKBOOK_PATH)

        prefix = "[DRY-RUN] " if dry_run else ""
        print(
            f"{prefix}Planilha: {WORKBOOK_PATH.name} | Aba: {SHEET_NAME} | "
            f"{new_rows} linha(s) nova(s) | {updated_values} valor(es) gravado(s)/atualizado(s)."
        )
        for entry in log_rows:
            print(entry)
        return {"new_rows": new_rows, "updated_values": updated_values}
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# ORQUESTRACAO
# ---------------------------------------------------------------------------
def coletar(dry_run=False):
    session, token = make_session_and_token()

    rotas = ROTAS
    if MODO_TODAS_ROTAS:
        rotas = [{"nome": "TODAS_AS_ROTAS"}]  # sem filtro de origem/destino
        log.info("Modo TODAS_AS_ROTAS ativo — recomendado JANELA='semanal'.")

    janelas = gerar_janelas(DATA_INICIO, DATA_FIM, JANELA)
    tarefas = [(rota, ini, fim) for rota in rotas for (ini, fim) in janelas]
    log.info(f"{len(tarefas)} requisicoes ({len(rotas)} rota(s) x {len(janelas)} janela(s), {JANELA}).")

    todos, truncados = [], []
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futs = [ex.submit(fetch_janela, session, token, r, i, f) for (r, i, f) in tarefas]
        for fut in as_completed(futs):
            res = fut.result()
            if res["ok"]:
                if res["n"]:
                    todos.extend(res["recs"])
                if res.get("truncado"):
                    truncados.append(f"{res['rota']} {res['janela']}")
            else:
                log.warning(f"Falha: {res['rota']} {res['janela']} -> {res['erro']}")

    if truncados:
        log.warning(f"{len(truncados)} janela(s) bateram no teto de {CAP_SERVIDOR} registros "
                    f"(dado pode estar truncado). Rode com JANELA='semanal'. Exemplos: {truncados[:5]}")

    if not todos:
        log.error("Nenhum registro retornado. Confira credenciais, indicador e filtros de rota.")
        return

    df = pd.DataFrame(todos).drop_duplicates().reset_index(drop=True)
    for col in ("Data", "data", "DataReferencia"):
        if col in df.columns:
            df = df.sort_values(col).reset_index(drop=True)
            break

    log.info(f"OK: {len(df)} linhas | colunas: {list(df.columns)}")
    update_workbook(df, dry_run=dry_run)


def main():
    parser = argparse.ArgumentParser(
        description="Coleta fretes do IMEA e atualiza a aba Preço Fretes da Transportes.xlsm."
    )
    parser.add_argument(
        "--dry-run", action="store_true",
        help="Coleta e mostra as alteracoes sem gravar na planilha.",
    )
    args = parser.parse_args()
    coletar(dry_run=args.dry_run)


if __name__ == "__main__":
    main()
